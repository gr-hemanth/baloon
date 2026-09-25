"""Student header metadata filler for worksheet documents.

Safely populates known student details (Name, Reg. No., Branch / Sec., Date)
into dedicated metadata tables or paragraphs without overwriting session,
topic, lecture, or instructional fields.
"""

from datetime import datetime
import logging
import re
from typing import Any, Dict, Optional

import docx
from docx.shared import Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

logger = logging.getLogger(__name__)

# Standard navy blue styling for filled student metadata
METADATA_COLOR_RGB = RGBColor(0x1F, 0x4E, 0x79)

# Default student details
DEFAULT_STUDENT_INFO: Dict[str, str] = {
    "name": "G R HEMANTH",
    "reg_no": "RA2511003011819",
    "branch": "COMPUTER SCIENCE AND ENGINEERING",
    "date": datetime.now().strftime("%d-%m-%Y"),
}

# Protected labels that must NEVER be overwritten
PROTECTED_HEADER_TERMS = [
    "session", "topic", "lecture", "course", "subject", "code", "faculty",
    "semester", "unit", "slo", "ps", "practice session", "module", "outcomes"
]


class StudentHeaderFiller:
    """Fills student identification details into document header tables and placeholders."""

    def __init__(self, student_info: Optional[Dict[str, str]] = None):
        self.student_info = {**DEFAULT_STUDENT_INFO, **(student_info or {})}

    def fill_header(self, doc: docx.Document) -> Dict[str, str]:
        """Inspect document tables and paragraphs to fill student details.

        Returns:
            Dictionary of fields populated (e.g. {'name': 'G R HEMANTH', ...}).
        """
        filled_fields: Dict[str, str] = {}

        # 1. Inspect top tables (first 3 tables typically hold metadata headers)
        for tbl_idx, table in enumerate(doc.tables[:3]):
            filled = self._fill_table_header(table)
            filled_fields.update(filled)

        # 2. Inspect top paragraphs (first 15 paragraphs)
        para_filled = self._fill_paragraph_header(doc.paragraphs[:15])
        filled_fields.update(para_filled)

        if filled_fields:
            logger.info("Populated student header fields: %s", list(filled_fields.keys()))
        return filled_fields

    def _fill_table_header(self, table: Table) -> Dict[str, str]:
        """Inspect table cells for Name, Reg. No., Branch, Date labels and populate adjacent blank cells."""
        filled: Dict[str, str] = {}

        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                raw_text = cell.text.strip()
                t_lower = raw_text.lower()

                # Check if cell is protected
                if any(term in t_lower for term in PROTECTED_HEADER_TERMS):
                    continue

                # Match NAME
                if "name" not in filled and self._is_name_label(t_lower):
                    target_cell = self._find_adjacent_blank_cell(row, c_idx)
                    if target_cell:
                        self._write_metadata_to_cell(target_cell, self.student_info["name"])
                        filled["name"] = self.student_info["name"]

                # Match REG. NO.
                elif "reg_no" not in filled and self._is_reg_no_label(t_lower):
                    target_cell = self._find_adjacent_blank_cell(row, c_idx)
                    if target_cell:
                        self._write_metadata_to_cell(target_cell, self.student_info["reg_no"])
                        filled["reg_no"] = self.student_info["reg_no"]

                # Match BRANCH / SEC.
                elif "branch" not in filled and self._is_branch_label(t_lower):
                    target_cell = self._find_adjacent_blank_cell(row, c_idx)
                    if target_cell:
                        self._write_metadata_to_cell(target_cell, self.student_info["branch"])
                        filled["branch"] = self.student_info["branch"]

                # Match DATE
                elif "date" not in filled and self._is_date_label(t_lower):
                    target_cell = self._find_adjacent_blank_cell(row, c_idx)
                    if target_cell:
                        self._write_metadata_to_cell(target_cell, self.student_info["date"])
                        filled["date"] = self.student_info["date"]

        return filled

    def _fill_paragraph_header(self, paragraphs: list) -> Dict[str, str]:
        """Inspect standalone header paragraphs with blank underscores or colons."""
        filled: Dict[str, str] = {}

        for p in paragraphs:
            text = p.text.strip()
            if not text:
                continue
            t_lower = text.lower()

            if any(term in t_lower for term in PROTECTED_HEADER_TERMS):
                continue

            # Check for patterns like "Name: ________" or "Name:"
            if "name" not in filled and self._is_name_label(t_lower):
                if "_" in text or text.endswith(":"):
                    self._replace_or_append_paragraph_value(p, "Name", self.student_info["name"])
                    filled["name"] = self.student_info["name"]

            elif "reg_no" not in filled and self._is_reg_no_label(t_lower):
                if "_" in text or text.endswith(":"):
                    self._replace_or_append_paragraph_value(p, "Reg. No.", self.student_info["reg_no"])
                    filled["reg_no"] = self.student_info["reg_no"]

            elif "branch" not in filled and self._is_branch_label(t_lower):
                if "_" in text or text.endswith(":"):
                    self._replace_or_append_paragraph_value(p, "Branch / Sec.", self.student_info["branch"])
                    filled["branch"] = self.student_info["branch"]

            elif "date" not in filled and self._is_date_label(t_lower):
                if "_" in text or text.endswith(":"):
                    self._replace_or_append_paragraph_value(p, "Date", self.student_info["date"])
                    filled["date"] = self.student_info["date"]

        return filled

    def _is_name_label(self, text: str) -> bool:
        return bool(re.match(r"^(?:student\s+)?name\s*[:\-–—]?$", text, re.IGNORECASE)) or (
            "name" in text and not any(p in text for p in ["course", "topic", "faculty", "father", "file"])
        )

    def _is_reg_no_label(self, text: str) -> bool:
        return bool(re.search(r"\b(?:reg(?:ister)?(?:\s+no|\.|\b)|roll\s+no|student\s+id)\b", text, re.IGNORECASE))

    def _is_branch_label(self, text: str) -> bool:
        return bool(re.search(r"\b(?:branch|dept|department|sec(?:tion)?)\b", text, re.IGNORECASE))

    def _is_date_label(self, text: str) -> bool:
        return bool(re.match(r"^date\s*[:\-–—]?$", text, re.IGNORECASE))

    def _find_adjacent_blank_cell(self, row: Any, col_idx: int) -> Optional[_Cell]:
        """Find the immediately following cell in the same row if it is empty."""
        if col_idx + 1 < len(row.cells):
            next_cell = row.cells[col_idx + 1]
            if not next_cell.text.strip():
                return next_cell
        return None

    def _write_metadata_to_cell(self, cell: _Cell, value: str) -> None:
        """Write formatted metadata value into target cell."""
        p = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
        p.text = ""
        run = p.add_run(value)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = METADATA_COLOR_RGB

    def _replace_or_append_paragraph_value(self, p: Paragraph, label: str, value: str) -> None:
        """Replace underline or append value to paragraph."""
        orig_text = p.text
        # Replace underline pattern if present
        if "_" in orig_text:
            new_text = re.sub(r"_{2,}", f" {value} ", orig_text).strip()
            p.text = new_text
            for run in p.runs:
                if value in run.text:
                    run.bold = True
                    run.font.color.rgb = METADATA_COLOR_RGB
        elif orig_text.rstrip().endswith(":"):
            run = p.add_run(f" {value}")
            run.bold = True
            run.font.color.rgb = METADATA_COLOR_RGB
