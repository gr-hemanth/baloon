"""Generic Answer-Target Resolution and Writing System for DOCX Worksheets.

Provides a robust, element-level target resolution engine that maps each parsed
question to its intended writable location (paragraph placeholder, table cell,
bordered box, text box, content control, or anchor) rather than guessing by order.
"""

import logging
import re
from enum import Enum
from typing import Any, Dict, List, Optional, Tuple

import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from packages.worksheets.answer_models import GeneratedAnswer
from packages.worksheets.exceptions import WorksheetFillingError
from packages.worksheets.models import ParsedQuestion, QuestionType, ResponseMode

logger = logging.getLogger(__name__)

# Standard styling colors for filled answers
ANSWER_COLOR_RGB = RGBColor(0x1F, 0x4E, 0x79)  # Professional navy blue
MUTED_GRAY_RGB = RGBColor(0x59, 0x59, 0x59)

# XML Namespaces used in DOCX structures
NAMESPACES = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "w14": "http://schemas.microsoft.com/office/word/2010/wordml",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
    "v": "urn:schemas-microsoft-com:vml",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
}


class AnswerTargetType(str, Enum):
    """Classification of writable answer target locations in DOCX documents."""
    PARAGRAPH_PLACEHOLDER = "PARAGRAPH_PLACEHOLDER"  # Existing paragraph containing "Answer:", "Ans:", etc.
    PARAGRAPH_BOX = "PARAGRAPH_BOX"                  # Bordered or designated empty paragraph for answer
    PARAGRAPH_EMPTY = "PARAGRAPH_EMPTY"              # Designated empty paragraph for answer
    TABLE_CELL_COLUMN = "TABLE_CELL_COLUMN"          # Dedicated answer column in same row
    TABLE_CELL_ROW = "TABLE_CELL_ROW"                # Dedicated answer row below question (alternating/merged)
    TABLE_CELL_BLANK = "TABLE_CELL_BLANK"            # Blank or bordered cell in same row
    TABLE_CELL_SPECIFIC = "TABLE_CELL_SPECIFIC"      # Exact resolved table cell matching target spec
    TABLE_CELL_APPEND = "TABLE_CELL_APPEND"          # Append inside question cell (when no dedicated cell)
    TEXT_BOX = "TEXT_BOX"                            # w:txbxContent element in DrawingML or VML
    CONTENT_CONTROL = "CONTENT_CONTROL"              # w:sdtContent structured document tag
    NEW_PARAGRAPH_ANCHOR = "NEW_PARAGRAPH_ANCHOR"    # Fallback: insert new paragraph after question block
    MULTI_TARGET = "MULTI_TARGET"                    # Multiple sub-targets (e.g. table activity or multi-cell table)


class AnswerTarget:
    """Resolved writable target for a single question."""

    def __init__(
        self,
        question_id: str,
        target_type: AnswerTargetType,
        paragraph_index: Optional[int] = None,
        paragraph_obj: Optional[Paragraph] = None,
        table_index: Optional[int] = None,
        row_index: Optional[int] = None,
        col_index: Optional[int] = None,
        cell_obj: Optional[_Cell] = None,
        txbx_element: Optional[Any] = None,
        sdt_element: Optional[Any] = None,
        placeholder_text: Optional[str] = None,
        is_explicit: bool = True,
        metadata: Optional[Dict[str, Any]] = None,
    ):
        self.question_id = question_id
        self.target_type = target_type
        self.paragraph_index = paragraph_index
        self.paragraph_obj = paragraph_obj
        self.table_index = table_index
        self.row_index = row_index
        self.col_index = col_index
        self.cell_obj = cell_obj
        self.txbx_element = txbx_element
        self.sdt_element = sdt_element
        self.placeholder_text = placeholder_text
        self.is_explicit = is_explicit
        self.metadata = metadata or {}

    def __repr__(self) -> str:
        return (
            f"<AnswerTarget q={self.question_id} type={self.target_type.value} "
            f"p_idx={self.paragraph_index} tbl={self.table_index} r={self.row_index} c={self.col_index} "
            f"explicit={self.is_explicit}>"
        )


def is_answer_placeholder_text(text: str) -> bool:
    """Check if text is an explicit answer prompt or placeholder."""
    t = text.strip()
    if not t:
        return False
    # Explicit answer labels
    if re.match(r"^(?:Answer|Ans|Solution|Response|Output)\s*[:\-–—]?\s*$", t, re.IGNORECASE):
        return True
    if re.match(r"^(?:Answer|Ans|Solution|Response)\s*[:\-–—]\s*(?:_{2,}|\.{3,}|\[.*?\])?\s*$", t, re.IGNORECASE):
        return True
    if re.match(r"^(?:Write|Enter|Type)\s+(?:your\s+)?answer\s*[:\-–—]?\s*$", t, re.IGNORECASE):
        return True
    if re.match(r"^(?:Your\s+Answer|Student\s+Answer)\s*[:\-–—]?\s*$", t, re.IGNORECASE):
        return True
    if re.match(r"^_{3,}$", t):
        return True
    if t in ["[Answer]", "[Answer Here]", "[Solution]", "[Write your answer here]"]:
        return True
    return False


def find_text_box_in_element(element: Any) -> Optional[Any]:
    """Locate any w:txbxContent element within a paragraph or cell element."""
    txbx = element.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}txbxContent")
    if txbx:
        return txbx[0]
    # Check alternate namespaces if any
    for child in element.iter():
        if child.tag.endswith("txbxContent"):
            return child
    return None


def find_content_control_in_element(element: Any) -> Optional[Any]:
    """Locate any w:sdtContent element within an element."""
    sdts = element.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sdtContent")
    if sdts:
        return sdts[0]
    for child in element.iter():
        if child.tag.endswith("sdtContent"):
            return child
    return None


def cell_has_borders(cell: _Cell) -> bool:
    """Check if a table cell has explicit border styling."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcBorders")
    return tcBorders is not None


class AnswerTargetResolver:
    """Resolves exactly one intended writable AnswerTarget for each question."""

    def resolve(
        self,
        doc: docx.Document,
        question: ParsedQuestion,
        all_questions: List[ParsedQuestion],
    ) -> AnswerTarget:
        """Resolve the exact writable answer target in doc for question.
        
        Args:
            doc: python-docx Document object.
            question: The question needing an answer target.
            all_questions: Full list of worksheet questions to determine boundaries.
            
        Returns:
            Resolved AnswerTarget.
            
        Raises:
            WorksheetFillingError: If no valid target can be resolved.
        """
        # 0. Check if question has explicit table cell targets defined (from parser)
        if question.targets and any(spec.target_type == "TABLE_CELL" for spec in question.targets):
            sub_targets: List[AnswerTarget] = []
            for spec in question.targets:
                if spec.target_type == "TABLE_CELL":
                    t_idx = spec.table_index
                    r_idx = spec.row_index
                    c_idx = spec.col_index
                    if t_idx is not None and t_idx < len(doc.tables):
                        table = doc.tables[t_idx]
                        if r_idx is not None and r_idx < len(table.rows):
                            row = table.rows[r_idx]
                            if c_idx is not None and c_idx < len(row.cells):
                                cell = row.cells[c_idx]
                                sub_targets.append(
                                    AnswerTarget(
                                        question_id=question.question_id,
                                        target_type=AnswerTargetType.TABLE_CELL_SPECIFIC,
                                        table_index=t_idx,
                                        row_index=r_idx,
                                        col_index=c_idx,
                                        cell_obj=cell,
                                        is_explicit=True,
                                        metadata={
                                            "target_id": spec.target_id,
                                            "semantic": spec.semantic,
                                            "column_header": spec.column_header,
                                            "row_label": spec.row_label,
                                            "expected_length": spec.expected_length,
                                        },
                                    )
                                )
            if sub_targets:
                if len(sub_targets) == 1:
                    return sub_targets[0]
                return AnswerTarget(
                    question_id=question.question_id,
                    target_type=AnswerTargetType.MULTI_TARGET,
                    is_explicit=True,
                    metadata={"sub_targets": sub_targets},
                )

        table_loc = question.source_location.get("table_index")
        row_loc = question.source_location.get("row_index")

        if table_loc is not None and row_loc is not None and table_loc < len(doc.tables):
            return self._resolve_table_target(doc.tables[table_loc], table_loc, row_loc, question)

        return self._resolve_paragraph_target(doc, question, all_questions)

    def _resolve_paragraph_target(
        self,
        doc: docx.Document,
        question: ParsedQuestion,
        all_questions: List[ParsedQuestion],
    ) -> AnswerTarget:
        """Resolve target for paragraph-based question."""
        q_p_idx = question.source_location.get("paragraph_index")
        placeholder_p_idx = question.source_location.get("answer_placeholder_paragraph_index")

        # 1. Direct match from parser metadata if valid
        if placeholder_p_idx is not None and placeholder_p_idx < len(doc.paragraphs):
            p = doc.paragraphs[placeholder_p_idx]
            return AnswerTarget(
                question_id=question.question_id,
                target_type=AnswerTargetType.PARAGRAPH_PLACEHOLDER,
                paragraph_index=placeholder_p_idx,
                paragraph_obj=p,
                placeholder_text=p.text.strip(),
                is_explicit=True,
            )

        # 2. Find question index if missing from source_location or shifted by earlier insertions
        is_valid_p = (
            q_p_idx is not None
            and q_p_idx < len(doc.paragraphs)
            and self._paragraph_matches_question(doc.paragraphs[q_p_idx], question)
        )
        if not is_valid_p:
            q_p_idx = self._find_question_paragraph_index(doc, question)

        if q_p_idx == -1:
            # If question paragraph couldn't be located at all, raise structured error
            raise WorksheetFillingError(
                f"Cannot locate question element in document for question ID '{question.question_id}' "
                f"(number={question.question_number}). No writable target could be resolved."
            )

        # Determine search window: from q_p_idx to next question's start index in current document
        next_q_idx = len(doc.paragraphs)
        q_order = question.source_order
        for other_q in all_questions:
            if other_q.source_order > q_order:
                other_p = self._find_question_paragraph_index(doc, other_q)
                if other_p != -1 and other_p > q_p_idx:
                    next_q_idx = min(next_q_idx, other_p)
                    break

        # 3. Scan window for explicit placeholder, textbox, content control, or bordered box
        last_content_idx = q_p_idx

        for idx in range(q_p_idx + 1, min(next_q_idx, len(doc.paragraphs))):
            p = doc.paragraphs[idx]
            text = p.text.strip()

            # Check for embedded text box in paragraph
            txbx = find_text_box_in_element(p._p)
            if txbx is not None:
                return AnswerTarget(
                    question_id=question.question_id,
                    target_type=AnswerTargetType.TEXT_BOX,
                    paragraph_index=idx,
                    paragraph_obj=p,
                    txbx_element=txbx,
                    is_explicit=True,
                )

            # Check for content control in paragraph
            sdt = find_content_control_in_element(p._p)
            if sdt is not None:
                return AnswerTarget(
                    question_id=question.question_id,
                    target_type=AnswerTargetType.CONTENT_CONTROL,
                    paragraph_index=idx,
                    paragraph_obj=p,
                    sdt_element=sdt,
                    is_explicit=True,
                )

            # Check for explicit answer placeholder text
            if is_answer_placeholder_text(text):
                return AnswerTarget(
                    question_id=question.question_id,
                    target_type=AnswerTargetType.PARAGRAPH_PLACEHOLDER,
                    paragraph_index=idx,
                    paragraph_obj=p,
                    placeholder_text=text,
                    is_explicit=True,
                )

            # Check if paragraph has borders designating an answer box
            pBdr = p._p.get_or_add_pPr().find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pBdr")
            if pBdr is not None and not text:
                return AnswerTarget(
                    question_id=question.question_id,
                    target_type=AnswerTargetType.PARAGRAPH_BOX,
                    paragraph_index=idx,
                    paragraph_obj=p,
                    is_explicit=True,
                )

            # Keep track of last option or description paragraph
            if text:
                last_content_idx = idx

        # 4. Check if there is an empty paragraph designated for answering following the question block
        if last_content_idx + 1 < min(next_q_idx, len(doc.paragraphs)):
            candidate_p = doc.paragraphs[last_content_idx + 1]
            if not candidate_p.text.strip():
                # Check if this is an isolated trailing empty paragraph at EOF in a document
                # where earlier questions had no designated empty paragraphs
                is_isolated_eof = (
                    (last_content_idx + 1 == len(doc.paragraphs) - 1)
                    and len(all_questions) > 1
                    and any(oq.source_order < question.source_order for oq in all_questions)
                    and not any(oq.targets for oq in all_questions if oq.question_id != question.question_id)
                )
                if not is_isolated_eof:
                    return AnswerTarget(
                        question_id=question.question_id,
                        target_type=AnswerTargetType.PARAGRAPH_EMPTY,
                        paragraph_index=last_content_idx + 1,
                        paragraph_obj=candidate_p,
                        is_explicit=True,
                    )

        # 5. Fallback: No explicit placeholder or blank line found; insert directly after last question block
        anchor_p = doc.paragraphs[last_content_idx]
        return AnswerTarget(
            question_id=question.question_id,
            target_type=AnswerTargetType.NEW_PARAGRAPH_ANCHOR,
            paragraph_index=last_content_idx,
            paragraph_obj=anchor_p,
            is_explicit=False,
        )

    def _resolve_table_target(
        self,
        table: Table,
        table_idx: int,
        row_idx: int,
        question: ParsedQuestion,
    ) -> AnswerTarget:
        """Resolve target inside a table."""
        if row_idx >= len(table.rows):
            raise WorksheetFillingError(
                f"Table row {row_idx} exceeds table row count ({len(table.rows)}) for question '{question.question_id}'."
            )

        row = table.rows[row_idx]

        # 1. Check if source_location identified an answer column
        ans_col = question.source_location.get("answer_col_index")
        if ans_col is not None and ans_col < len(row.cells):
            return AnswerTarget(
                question_id=question.question_id,
                target_type=AnswerTargetType.TABLE_CELL_COLUMN,
                table_index=table_idx,
                row_index=row_idx,
                col_index=ans_col,
                cell_obj=row.cells[ans_col],
                is_explicit=True,
            )

        # 2. Inspect table header for dedicated Answer / Solution column
        header_cells = [c.text.strip().lower() for c in table.rows[0].cells]
        for idx, h in enumerate(header_cells):
            if any(term in h for term in ["answer", "solution", "response", "output", "result"]):
                if idx < len(row.cells):
                    return AnswerTarget(
                        question_id=question.question_id,
                        target_type=AnswerTargetType.TABLE_CELL_COLUMN,
                        table_index=table_idx,
                        row_index=row_idx,
                        col_index=idx,
                        cell_obj=row.cells[idx],
                        is_explicit=True,
                    )

        # 3. Check for dedicated answer row directly below (row-alternating table pattern)
        if row_idx + 1 < len(table.rows):
            next_row = table.rows[row_idx + 1]
            # If next row has a cell starting with "Answer:" or is empty with borders
            next_c0 = next_row.cells[0]
            next_text = next_c0.text.strip()
            if is_answer_placeholder_text(next_text) or (not next_text and cell_has_borders(next_c0)):
                return AnswerTarget(
                    question_id=question.question_id,
                    target_type=AnswerTargetType.TABLE_CELL_ROW,
                    table_index=table_idx,
                    row_index=row_idx + 1,
                    col_index=0,
                    cell_obj=next_c0,
                    placeholder_text=next_text,
                    is_explicit=True,
                )

        # 4. Check for blank or bordered cell in the same row (other than question cell)
        q_col = question.source_location.get("col_index", 0)
        for c_idx, cell in enumerate(row.cells):
            if c_idx != q_col:
                c_text = cell.text.strip()
                if not c_text or is_answer_placeholder_text(c_text):
                    return AnswerTarget(
                        question_id=question.question_id,
                        target_type=AnswerTargetType.TABLE_CELL_BLANK,
                        table_index=table_idx,
                        row_index=row_idx,
                        col_index=c_idx,
                        cell_obj=cell,
                        placeholder_text=c_text,
                        is_explicit=True,
                    )

        # 5. Check if question cell itself contains a text box
        target_cell = row.cells[min(q_col, len(row.cells) - 1)]
        txbx = find_text_box_in_element(target_cell._tc)
        if txbx is not None:
            return AnswerTarget(
                question_id=question.question_id,
                target_type=AnswerTargetType.TEXT_BOX,
                table_index=table_idx,
                row_index=row_idx,
                col_index=q_col,
                cell_obj=target_cell,
                txbx_element=txbx,
                is_explicit=True,
            )

        # 6. Fallback: Append inside question cell
        return AnswerTarget(
            question_id=question.question_id,
            target_type=AnswerTargetType.TABLE_CELL_APPEND,
            table_index=table_idx,
            row_index=row_idx,
            col_index=q_col,
            cell_obj=target_cell,
            is_explicit=False,
        )

    def _paragraph_matches_question(self, p: Paragraph, question: ParsedQuestion) -> bool:
        """Check if a paragraph element in doc matches the question."""
        p_text = p.text.strip().lower()
        if not p_text:
            return False
        if p_text.startswith("answer:") or p_text.startswith("ans:") or p_text.startswith("solution:"):
            return False

        q_lines = [line.strip().lower() for line in question.question_text.splitlines() if line.strip()]
        if not q_lines:
            return False
        q_clean = q_lines[0][:50]

        if q_clean in p_text or p_text in q_clean:
            return True

        q_num = question.question_number
        if q_num:
            patterns = [
                f"{q_num}.", f"{q_num})", f"q{q_num}.", f"q{q_num}:",
                f"activity {q_num}", f"activity {q_num}:"
            ]
            if any(p_text.startswith(pat) or f" {pat}" in p_text for pat in patterns):
                if len(q_clean) < 15 or any(word in p_text for word in q_clean.split()[:4]):
                    return True

        return False

    def _find_question_paragraph_index(self, doc: docx.Document, question: ParsedQuestion) -> int:
        """Find paragraph index matching question text and number in current document state."""
        for idx, p in enumerate(doc.paragraphs):
            if self._paragraph_matches_question(p, question):
                return idx

        return -1


class TargetWriter:
    """Writes answers cleanly into resolved AnswerTargets, preserving styles and structure."""

    def write(
        self,
        doc: docx.Document,
        target: AnswerTarget,
        question: ParsedQuestion,
        answer: GeneratedAnswer,
    ) -> None:
        """Write formatted answer into target location.
        
        Args:
            doc: Document being edited.
            target: Resolved AnswerTarget.
            question: The ParsedQuestion being answered.
            answer: Generated solution.
        """
        is_code = self._is_code_answer(question, answer)
        formatted_text = self._format_answer_text(question, answer)
        if is_code:
            lines = [line.rstrip() for line in formatted_text.splitlines()]
            while lines and not lines[0].strip():
                lines.pop(0)
            while lines and not lines[-1].strip():
                lines.pop()
        else:
            lines = [line.strip() for line in formatted_text.splitlines() if line.strip()]

        if target.target_type == AnswerTargetType.MULTI_TARGET:
            for sub_target in target.metadata.get("sub_targets", []):
                self._write_sub_target(doc, sub_target, question, answer)

        elif target.target_type == AnswerTargetType.PARAGRAPH_EMPTY:
            self._write_paragraph_empty(target.paragraph_obj, lines, question, is_code=is_code)

        elif target.target_type == AnswerTargetType.TABLE_CELL_SPECIFIC:
            self._write_table_cell_specific(target.cell_obj, lines, target.metadata, is_code=is_code)

        elif target.target_type == AnswerTargetType.PARAGRAPH_PLACEHOLDER:
            self._write_paragraph_placeholder(target.paragraph_obj, lines, target.placeholder_text, is_code=is_code)

        elif target.target_type == AnswerTargetType.PARAGRAPH_BOX:
            self._write_paragraph_box(target.paragraph_obj, lines, is_code=is_code)

        elif target.target_type in (
            AnswerTargetType.TABLE_CELL_COLUMN,
            AnswerTargetType.TABLE_CELL_ROW,
            AnswerTargetType.TABLE_CELL_BLANK,
        ):
            self._write_table_cell(target.cell_obj, lines, target.placeholder_text, is_code=is_code)

        elif target.target_type == AnswerTargetType.TABLE_CELL_APPEND:
            self._write_table_cell_append(target.cell_obj, lines, is_code=is_code)

        elif target.target_type == AnswerTargetType.TEXT_BOX:
            self._write_text_box(doc, target.txbx_element, lines)

        elif target.target_type == AnswerTargetType.CONTENT_CONTROL:
            self._write_content_control(doc, target.sdt_element, lines)

        elif target.target_type == AnswerTargetType.NEW_PARAGRAPH_ANCHOR:
            self._write_new_paragraph_anchor(target.paragraph_obj, lines, is_code=is_code)

        else:
            raise WorksheetFillingError(f"Unsupported target type: {target.target_type}")

        # If MCQ, bold the selected option in the question block if options were separate paragraphs
        if question.question_type == QuestionType.MCQ and answer.selected_option:
            self._highlight_mcq_option(doc, question, answer.selected_option)

    def _is_code_answer(self, question: ParsedQuestion, answer: GeneratedAnswer) -> bool:
        """Determine whether answer should be rendered with code styling."""
        resp_mode = getattr(answer, "response_mode", None) or getattr(question, "response_mode", None)
        if resp_mode in (ResponseMode.CODE, ResponseMode.CODE_AND_EXPLANATION):
            return True
        return question.question_type == QuestionType.CODE

    def _style_line_run(
        self,
        paragraph: Paragraph,
        run,
        line: str,
        is_code: bool,
        is_code_line: bool,
    ) -> None:
        """Apply font, size, and paragraph spacing according to code or text mode."""
        if is_code_line:
            run.font.name = "Consolas"
            run.font.size = Pt(9.0)
            paragraph.paragraph_format.space_after = Pt(1.5)
            paragraph.paragraph_format.line_spacing = 1.15
        else:
            run.font.size = Pt(10.0 if is_code else 10.5)
        run.font.color.rgb = ANSWER_COLOR_RGB

    def _format_answer_text(self, question: ParsedQuestion, answer: GeneratedAnswer) -> str:
        """Format answer text cleanly based on classification."""
        is_code = self._is_code_answer(question, answer)
        text = answer.answer_text.strip()

        # Remove redundant leading "Answer:" or "Ans:" from answer engine output if present
        text = re.sub(r"^(?:Answer|Ans|Solution)\s*[:\-–—]\s*", "", text, flags=re.IGNORECASE).strip()

        if question.question_type == QuestionType.MCQ and answer.selected_option:
            opt_key = answer.selected_option.upper()
            # If answer text doesn't already include option letter prefix
            if not text.upper().startswith(opt_key):
                return f"{opt_key}. {text}"
        return self._clean_humanized_text(text, is_code=is_code)

    @staticmethod
    def _clear_paragraph(p: Paragraph) -> None:
        """Clear all runs and text from a paragraph without leaving an unstyled empty run."""
        p.text = ""
        while p.runs:
            r = p.runs[0]
            if r._r.getparent() is not None:
                r._r.getparent().remove(r._r)
            else:
                break

    def _clean_humanized_text(self, text: str, is_code: bool = False) -> str:
        """Strip markdown markers (###, **, *), AI buzzwords, and redundant prefixes."""
        t = text.strip()
        # Clean markdown code fences if present
        t = re.sub(r"^```[a-zA-Z0-9_-]*\s*\n?", "", t, flags=re.MULTILINE)
        t = re.sub(r"\n?```\s*$", "", t, flags=re.MULTILINE)
        if not is_code:
            t = re.sub(r"\*\*([^*]+)\*\*", r"\1", t)
            t = re.sub(r"^#{1,6}\s*", "", t, flags=re.MULTILINE)
            t = re.sub(r"^\s*[\*\-•]\s+", "", t, flags=re.MULTILINE)
        t = re.sub(r"^(?:Answer|Ans|Solution)\s*[:\-–—]\s*", "", t, flags=re.IGNORECASE).strip()
        return t

    def _write_sub_target(
        self,
        doc: docx.Document,
        sub_target: AnswerTarget,
        question: ParsedQuestion,
        answer: GeneratedAnswer,
    ) -> None:
        """Write into an individual sub-target within a multi-target question."""
        target_id = sub_target.metadata.get("target_id", "")
        text = ""
        if answer.target_answers and target_id in answer.target_answers:
            text = answer.target_answers[target_id]
        elif answer.target_answers:
            for k, v in answer.target_answers.items():
                if k in target_id or target_id.endswith(k):
                    text = v
                    break
        if not text:
            text = answer.answer_text

        is_code = self._is_code_answer(question, answer)
        text = self._clean_humanized_text(text, is_code=is_code)
        if is_code:
            lines = [line.rstrip() for line in text.splitlines()]
            while lines and not lines[0].strip():
                lines.pop(0)
            while lines and not lines[-1].strip():
                lines.pop()
        else:
            lines = [line.strip() for line in text.splitlines() if line.strip()]

        if sub_target.target_type in (
            AnswerTargetType.TABLE_CELL_SPECIFIC,
            AnswerTargetType.TABLE_CELL_COLUMN,
            AnswerTargetType.TABLE_CELL_BLANK,
        ):
            self._write_table_cell_specific(sub_target.cell_obj, lines, sub_target.metadata, is_code=is_code)
        elif sub_target.target_type == AnswerTargetType.PARAGRAPH_EMPTY:
            self._write_paragraph_empty(sub_target.paragraph_obj, lines, question, is_code=is_code)
        elif sub_target.target_type == AnswerTargetType.PARAGRAPH_PLACEHOLDER:
            self._write_paragraph_placeholder(sub_target.paragraph_obj, lines, sub_target.placeholder_text, is_code=is_code)

    def _write_table_cell_specific(
        self,
        cell: _Cell,
        lines: List[str],
        metadata: Dict[str, Any],
        is_code: bool = False,
    ) -> None:
        """Write concise answer cleanly into a specific table cell without any 'Answer:' prefix."""
        p = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
        self._clear_paragraph(p)
        if not lines:
            return

        in_explanation = False
        first = True
        for line in lines:
            if not line.strip() and not is_code:
                continue
            if line.strip().lower().startswith("explanation:"):
                in_explanation = True

            is_code_line = is_code and not in_explanation
            if first:
                cur_p = p
                txt = line
                if not is_code:
                    txt = re.sub(r"^(?:Answer|Ans|Solution)\s*[:\-–—]\s*", "", txt, flags=re.IGNORECASE).strip()
                r = cur_p.add_run(txt)
                self._style_line_run(cur_p, r, txt, is_code, is_code_line)
                first = False
            else:
                cur_p = cell.add_paragraph()
                r2 = cur_p.add_run(line if is_code else line.strip())
                self._style_line_run(cur_p, r2, line, is_code, is_code_line)

    def _write_paragraph_empty(
        self,
        p: Paragraph,
        lines: List[str],
        question: ParsedQuestion,
        is_code: bool = False,
    ) -> None:
        """Write answer cleanly into designated empty paragraph(s) without spurious headers."""
        self._clear_paragraph(p)
        if not lines:
            return

        first = True
        current_p = p
        parent_doc = p._parent
        in_explanation = False

        for line in lines:
            if not line.strip() and not is_code:
                continue
            if line.strip().lower().startswith("explanation:"):
                in_explanation = True

            is_code_line = is_code and not in_explanation

            if first:
                run = current_p.add_run(line if is_code else line.strip())
                self._style_line_run(current_p, run, line, is_code, is_code_line)
                first = False
            else:
                # Check if immediate next sibling is an existing empty paragraph we can reuse
                next_elm = current_p._p.getnext()
                is_reusable_empty_p = False
                if next_elm is not None and next_elm.tag.endswith("p"):
                    t_nodes = next_elm.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t")
                    if not any(t.text and t.text.strip() for t in t_nodes):
                        is_reusable_empty_p = True

                if is_reusable_empty_p:
                    cont_p = Paragraph(next_elm, parent_doc)
                    self._clear_paragraph(cont_p)
                    run = cont_p.add_run(line if is_code else line.strip())
                    self._style_line_run(cont_p, run, line, is_code, is_code_line)
                    current_p = cont_p
                else:
                    new_p_elm = current_p._p.getparent()._new_p()
                    current_p._p.addnext(new_p_elm)
                    cont_p = Paragraph(new_p_elm, parent_doc)
                    run = cont_p.add_run(line if is_code else line.strip())
                    self._style_line_run(cont_p, run, line, is_code, is_code_line)
                    current_p = cont_p

    def _write_paragraph_placeholder(
        self,
        p: Paragraph,
        lines: List[str],
        placeholder_text: Optional[str],
        is_code: bool = False,
    ) -> None:
        """Write into an existing placeholder paragraph (e.g. 'Answer: ')."""
        p_text = (placeholder_text or p.text).strip()

        # Clear existing text / runs
        self._clear_paragraph(p)

        # Check if the placeholder contained a label like "Answer:"
        has_label = bool(re.match(r"^(?:Answer|Ans|Solution|Response|Output)\s*[:\-–—]?", p_text, re.IGNORECASE))
        label_str = "Answer: " if has_label else ""

        if label_str:
            lbl_run = p.add_run(label_str)
            lbl_run.bold = True
            lbl_run.font.color.rgb = ANSWER_COLOR_RGB
            lbl_run.font.size = Pt(10.5)

        in_explanation = False
        if lines:
            line0 = lines[0]
            if line0.strip().lower().startswith("explanation:"):
                in_explanation = True
            is_code_line0 = is_code and not in_explanation
            txt_run = p.add_run(line0 if is_code else line0.strip())
            self._style_line_run(p, txt_run, line0, is_code, is_code_line0)

        # For multi-line responses, insert continuation paragraphs immediately following
        current_anchor = p
        parent_doc = p._parent
        for line in lines[1:]:
            if not line.strip() and not is_code:
                continue
            if line.strip().lower().startswith("explanation:"):
                in_explanation = True
            is_code_line = is_code and not in_explanation

            new_p_elm = current_anchor._p.getparent()._new_p()
            current_anchor._p.addnext(new_p_elm)
            cont_p = Paragraph(new_p_elm, parent_doc)
            r = cont_p.add_run(line if is_code else line.strip())
            if not is_code and re.match(r"^\d+[\.\)]\s+[A-Za-z]", line.strip()):
                r.bold = True
            self._style_line_run(cont_p, r, line, is_code, is_code_line)
            current_anchor = cont_p

    def _write_paragraph_box(self, p: Paragraph, lines: List[str], is_code: bool = False) -> None:
        """Write into an empty bordered paragraph box."""
        self._clear_paragraph(p)
        lbl_run = p.add_run("Answer: ")
        lbl_run.bold = True
        lbl_run.font.color.rgb = ANSWER_COLOR_RGB
        lbl_run.font.size = Pt(10.5)

        in_explanation = False
        if lines:
            line0 = lines[0]
            if line0.strip().lower().startswith("explanation:"):
                in_explanation = True
            is_code_line0 = is_code and not in_explanation
            txt_run = p.add_run(line0 if is_code else line0.strip())
            self._style_line_run(p, txt_run, line0, is_code, is_code_line0)

        current_anchor = p
        parent_doc = p._parent
        for line in lines[1:]:
            if not line.strip() and not is_code:
                continue
            if line.strip().lower().startswith("explanation:"):
                in_explanation = True
            is_code_line = is_code and not in_explanation

            new_p_elm = current_anchor._p.getparent()._new_p()
            current_anchor._p.addnext(new_p_elm)
            cont_p = Paragraph(new_p_elm, parent_doc)
            r = cont_p.add_run(line if is_code else line.strip())
            self._style_line_run(cont_p, r, line, is_code, is_code_line)
            current_anchor = cont_p

    def _write_table_cell(
        self,
        cell: _Cell,
        lines: List[str],
        placeholder_text: Optional[str],
        is_code: bool = False,
    ) -> None:
        """Write into a dedicated table cell."""
        p = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
        self._clear_paragraph(p)

        # Check if cell had an "Answer:" label
        if placeholder_text and re.match(r"^(?:Answer|Ans|Solution)\s*[:\-–—]?", placeholder_text, re.IGNORECASE):
            lbl_run = p.add_run("Answer: ")
            lbl_run.bold = True
            lbl_run.font.color.rgb = ANSWER_COLOR_RGB
            lbl_run.font.size = Pt(10)

        in_explanation = False
        if lines:
            line0 = lines[0]
            if line0.strip().lower().startswith("explanation:"):
                in_explanation = True
            is_code_line0 = is_code and not in_explanation
            txt_run = p.add_run(line0 if is_code else line0.strip())
            self._style_line_run(p, txt_run, line0, is_code, is_code_line0)

        # Multi-paragraph inside the same cell
        for line in lines[1:]:
            if not line.strip() and not is_code:
                continue
            if line.strip().lower().startswith("explanation:"):
                in_explanation = True
            is_code_line = is_code and not in_explanation

            new_p = cell.add_paragraph()
            r = new_p.add_run(line if is_code else line.strip())
            if not is_code and re.match(r"^\d+[\.\)]\s+[A-Za-z]", line.strip()):
                r.bold = True
            self._style_line_run(new_p, r, line, is_code, is_code_line)

    def _write_table_cell_append(self, cell: _Cell, lines: List[str], is_code: bool = False) -> None:
        """Append answer inside question cell when no separate cell exists."""
        ans_p = cell.add_paragraph()
        lbl_run = ans_p.add_run("Answer: ")
        lbl_run.bold = True
        lbl_run.font.color.rgb = ANSWER_COLOR_RGB
        lbl_run.font.size = Pt(10)

        in_explanation = False
        if lines:
            line0 = lines[0]
            if line0.strip().lower().startswith("explanation:"):
                in_explanation = True
            is_code_line0 = is_code and not in_explanation
            txt_run = ans_p.add_run(line0 if is_code else line0.strip())
            self._style_line_run(ans_p, txt_run, line0, is_code, is_code_line0)

        for line in lines[1:]:
            if not line.strip() and not is_code:
                continue
            if line.strip().lower().startswith("explanation:"):
                in_explanation = True
            is_code_line = is_code and not in_explanation

            new_p = cell.add_paragraph()
            r = new_p.add_run(line if is_code else line.strip())
            self._style_line_run(new_p, r, line, is_code, is_code_line)

    def _write_text_box(self, doc: docx.Document, txbx_element: Any, lines: List[str]) -> None:
        """Write answer into w:txbxContent element in DrawingML or VML."""
        p_elems = txbx_element.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p")
        if not p_elems:
            new_p_elm = doc._body._element._new_p()
            txbx_element.append(new_p_elm)
            p_elems = [new_p_elm]

        first_p = Paragraph(p_elems[0], doc)
        first_p.text = ""
        lbl_run = first_p.add_run("Answer: ")
        lbl_run.bold = True
        lbl_run.font.color.rgb = ANSWER_COLOR_RGB
        lbl_run.font.size = Pt(10)

        if lines:
            txt_run = first_p.add_run(lines[0])
            txt_run.font.color.rgb = ANSWER_COLOR_RGB
            txt_run.font.size = Pt(10)

        for line in lines[1:]:
            if not line.strip():
                continue
            new_p_elm = doc._body._element._new_p()
            txbx_element.append(new_p_elm)
            cont_p = Paragraph(new_p_elm, doc)
            r = cont_p.add_run(line)
            r.font.color.rgb = ANSWER_COLOR_RGB
            r.font.size = Pt(10)

    def _write_content_control(self, doc: docx.Document, sdt_element: Any, lines: List[str]) -> None:
        """Write answer into w:sdtContent structured document tag."""
        p_elems = sdt_element.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p")
        if not p_elems:
            new_p_elm = doc._body._element._new_p()
            sdt_element.append(new_p_elm)
            p_elems = [new_p_elm]

        first_p = Paragraph(p_elems[0], doc)
        first_p.text = ""
        lbl_run = first_p.add_run("Answer: ")
        lbl_run.bold = True
        lbl_run.font.color.rgb = ANSWER_COLOR_RGB
        lbl_run.font.size = Pt(10.5)

        if lines:
            txt_run = first_p.add_run(lines[0])
            txt_run.font.color.rgb = ANSWER_COLOR_RGB
            txt_run.font.size = Pt(10.5)

        current_anchor = first_p
        for line in lines[1:]:
            if not line.strip():
                continue
            new_p_elm = doc._body._element._new_p()
            sdt_element.append(new_p_elm)
            cont_p = Paragraph(new_p_elm, doc)
            r = cont_p.add_run(line)
            r.font.color.rgb = ANSWER_COLOR_RGB
            r.font.size = Pt(10.5)

    def _write_new_paragraph_anchor(self, anchor_p: Paragraph, lines: List[str], is_code: bool = False) -> None:
        """Insert new paragraph immediately following anchor paragraph."""
        parent_doc = anchor_p._parent
        current_anchor = anchor_p

        new_p_elm = anchor_p._p.getparent()._new_p()
        current_anchor._p.addnext(new_p_elm)
        ans_p = Paragraph(new_p_elm, parent_doc)

        label_run = ans_p.add_run("Answer: ")
        label_run.bold = True
        label_run.font.color.rgb = ANSWER_COLOR_RGB
        label_run.font.size = Pt(10.5)

        in_explanation = False
        if lines:
            line0 = lines[0]
            if line0.strip().lower().startswith("explanation:"):
                in_explanation = True
            is_code_line0 = is_code and not in_explanation
            text_run = ans_p.add_run(line0 if is_code else line0.strip())
            self._style_line_run(ans_p, text_run, line0, is_code, is_code_line0)
        current_anchor = ans_p

        for line in lines[1:]:
            if not line.strip() and not is_code:
                continue
            if line.strip().lower().startswith("explanation:"):
                in_explanation = True
            is_code_line = is_code and not in_explanation

            line_elm = current_anchor._p.getparent()._new_p()
            current_anchor._p.addnext(line_elm)
            line_p = Paragraph(line_elm, parent_doc)

            run = line_p.add_run(line if is_code else line.strip())
            if not is_code and re.match(r"^\d+\.\s+[A-Za-z]", line.strip()):
                run.bold = True
            self._style_line_run(line_p, run, line, is_code, is_code_line)
            current_anchor = line_p

    def _highlight_mcq_option(self, doc: docx.Document, question: ParsedQuestion, selected_option: str) -> None:
        """Bold and highlight matching MCQ option in document paragraphs if separated."""
        q_p_idx = question.source_location.get("paragraph_index")
        if q_p_idx is None or q_p_idx >= len(doc.paragraphs):
            return

        opt_prefix = selected_option.upper()
        # Search paragraphs following question for the option
        for idx in range(q_p_idx + 1, min(q_p_idx + 8, len(doc.paragraphs))):
            p = doc.paragraphs[idx]
            t = p.text.strip().upper()
            if re.match(rf"^(?:\({opt_prefix}\)|\[{opt_prefix}\]|{opt_prefix}[\.\)])\s+", t):
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = ANSWER_COLOR_RGB
                break
