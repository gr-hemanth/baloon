"""Unit and integration tests for Generic Worksheet Parsing Engine.

Tests:
1. Pure MCQ worksheet parsing and options extraction.
2. One-word / fill-in-the-blank worksheet parsing.
3. Short-answer question classification and marks parsing.
4. Long-answer / design thinking / simulation parsing.
5. Mixed worksheet parsing across distinct sections.
6. Table-based worksheet parsing.
7. Safe round-trip test (zero modification of original document, deterministic ordering).
8. Real 1011.docx structure inspection (if artifact is present).
9. PDF parsing abstraction and dispatch.
10. Processor interface verification.
"""

import hashlib
import shutil
from pathlib import Path
from unittest.mock import MagicMock, patch

import docx
import pytest

from packages.worksheets.classifier import QuestionClassifier
from packages.worksheets.docx_parser import DocxWorksheetParser
from packages.worksheets.models import (
    ParsedQuestion,
    ParsedWorksheet,
    QuestionOption,
    QuestionType,
)
from packages.worksheets.parser import WorksheetParser
from packages.worksheets.pdf_parser import PdfWorksheetParser
from packages.worksheets.processor import DefaultWorksheetProcessor


# ---------------------------------------------------------------------------
# Helpers to generate synthetic DOCX files
# ---------------------------------------------------------------------------

def _create_synthetic_mcq_docx(path: Path) -> Path:
    doc = docx.Document()
    doc.add_heading("Part A - Multiple Choice Questions", level=1)

    # Q1: Standard with separate options
    doc.add_paragraph("1. What does CPU stand for?")
    doc.add_paragraph("A. Central Processing Unit")
    doc.add_paragraph("B. Computer Personal Unit")
    doc.add_paragraph("C. Central Performance Utility")
    doc.add_paragraph("D. Core Processing Unit")

    # Q2: Standard with parentheses
    doc.add_paragraph("2. Which data structure operates on LIFO principle?")
    doc.add_paragraph("(a) Queue")
    doc.add_paragraph("(b) Stack")
    doc.add_paragraph("(c) Tree")
    doc.add_paragraph("(d) Graph")

    # Q3: Question with inline options
    doc.add_paragraph(
        "3. Which protocol is used for secure web browsing? "
        "A. HTTP B. HTTPS C. FTP D. SSH"
    )

    # Q4: Question with Q prefix
    doc.add_paragraph("Q4. Which layer of the OSI model is responsible for routing?")
    doc.add_paragraph("A. Physical layer")
    doc.add_paragraph("B. Data Link layer")
    doc.add_paragraph("C. Network layer")
    doc.add_paragraph("D. Transport layer")

    doc.save(str(path))
    return path


def _create_synthetic_one_word_docx(path: Path) -> Path:
    doc = docx.Document()
    doc.add_heading("Part A - Fill in the Blanks", level=1)

    doc.add_paragraph("1. The process of converting source code into machine code is called ________.")
    doc.add_paragraph("2. Python is an interpreted programming language. (True or False)")
    doc.add_paragraph("3. Fill in the blank: The default port for HTTP is ( ).")
    doc.add_paragraph("4. Name the following: A software design pattern that ensures a class has only one instance.")

    doc.save(str(path))
    return path


def _create_synthetic_short_answer_docx(path: Path) -> Path:
    doc = docx.Document()
    doc.add_heading("Part B - Short Answer Questions", level=1)

    doc.add_paragraph("1. Define polymorphism in object-oriented programming. [2 Marks]")
    doc.add_paragraph("2. Differentiate between process and thread. [2 Marks]")
    doc.add_paragraph("3. List any three advantages of Agile methodology.")
    doc.add_paragraph("4. Briefly explain the function of an operating system kernel. (3M)")

    doc.save(str(path))
    return path


def _create_synthetic_long_answer_docx(path: Path) -> Path:
    doc = docx.Document()
    doc.add_heading("Part C - Long Answer Questions", level=1)

    doc.add_paragraph(
        "1. Explain the architectural differences between monolithic and microservices "
        "architectures in detail. [10 Marks]"
    )
    doc.add_paragraph(
        "2. Case Study: Design a distributed caching layer for a high-traffic e-commerce "
        "application and discuss fault-tolerance strategies. [16 Marks]"
    )
    doc.add_paragraph(
        "3. Role-Play Simulation: Develop a comprehensive project proposal for cloud migration."
    )

    doc.save(str(path))
    return path


def _create_synthetic_mixed_docx(path: Path) -> Path:
    doc = docx.Document()
    # Header metadata
    doc.add_paragraph("Unit II – Session 10 – SLO 2")
    doc.add_paragraph("Course Code: 21CSC303J")

    # Section 1: MCQ
    doc.add_heading("Part A - Objective Section", level=1)
    doc.add_paragraph("1. What is the time complexity of binary search?")
    doc.add_paragraph("A. O(1)")
    doc.add_paragraph("B. O(n)")
    doc.add_paragraph("C. O(log n)")
    doc.add_paragraph("D. O(n^2)")
    doc.add_paragraph("2. A function that calls itself is known as ________ function.")

    # Section 2: Short Answer
    doc.add_heading("Part B - Conceptual Questions", level=1)
    doc.add_paragraph("3. State Amdahl's Law. [2 Marks]")

    # Section 3: Long Answer
    doc.add_heading("Part C - Detailed Problems", level=1)
    doc.add_paragraph(
        "4. Explain in detail the Raft consensus algorithm used in distributed systems. [10 Marks]"
    )

    doc.save(str(path))
    return path


def _create_synthetic_table_docx(path: Path) -> Path:
    doc = docx.Document()
    doc.add_heading("Section A - Table Questions", level=1)

    table = doc.add_table(rows=4, cols=3)
    # Header
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "S.No"
    hdr_cells[1].text = "Question Description"
    hdr_cells[2].text = "Marks"

    # Row 1
    r1 = table.rows[1].cells
    r1[0].text = "1"
    r1[1].text = "Define software configuration management."
    r1[2].text = "2"

    # Row 2
    r2 = table.rows[2].cells
    r2[0].text = "2"
    r2[1].text = "Explain the Spiral Model with a neat diagram."
    r2[2].text = "10"

    # Row 3 (MCQ inside table)
    r3 = table.rows[3].cells
    r3[0].text = "3"
    r3[1].text = "Which is a valid SDLC phase? A. Requirements B. Sleeping C. Wandering D. Playing"
    r3[2].text = "1"

    doc.save(str(path))
    return path


# ---------------------------------------------------------------------------
# Test Cases
# ---------------------------------------------------------------------------

def test_synthetic_mcq_worksheet(tmp_path: Path):
    """Verify parsing and option extraction for multiple choice questions."""
    docx_file = _create_synthetic_mcq_docx(tmp_path / "mcq_test.docx")
    parser = DocxWorksheetParser()
    parsed = parser.parse(docx_file)

    assert parsed.file_format == "docx"
    assert len(parsed.questions) == 4
    for q in parsed.questions:
        assert q.question_type == QuestionType.MCQ
        assert len(q.options) == 4

    # Check first question options
    q1 = parsed.questions[0]
    assert q1.question_number == "1"
    assert "CPU stand for" in q1.question_text
    assert q1.options[0].key == "A"
    assert "Central Processing Unit" in q1.options[0].text
    assert q1.options[3].key == "D"

    # Check inline options on Q3
    q3 = parsed.questions[2]
    assert q3.question_number == "3"
    assert len(q3.options) == 4
    assert q3.options[1].key == "B"
    assert q3.options[1].text == "HTTPS"


def test_synthetic_one_word_worksheet(tmp_path: Path):
    """Verify classification of fill-in-the-blank, one-word, and boolean questions."""
    docx_file = _create_synthetic_one_word_docx(tmp_path / "one_word_test.docx")
    parser = DocxWorksheetParser()
    parsed = parser.parse(docx_file)

    assert len(parsed.questions) == 4
    for q in parsed.questions:
        assert q.question_type == QuestionType.ONE_WORD

    assert parsed.questions[0].question_number == "1"
    assert "converting source code" in parsed.questions[0].question_text
    assert parsed.questions[1].question_number == "2"
    assert "interpreted programming" in parsed.questions[1].question_text


def test_synthetic_short_answer_worksheet(tmp_path: Path):
    """Verify classification of short answer questions and marks extraction."""
    docx_file = _create_synthetic_short_answer_docx(tmp_path / "short_answer_test.docx")
    parser = DocxWorksheetParser()
    parsed = parser.parse(docx_file)

    assert len(parsed.questions) == 4
    for q in parsed.questions:
        assert q.question_type == QuestionType.SHORT_ANSWER

    assert parsed.questions[0].marks == 2.0
    assert parsed.questions[1].marks == 2.0
    assert parsed.questions[3].marks == 3.0


def test_synthetic_long_answer_worksheet(tmp_path: Path):
    """Verify classification of long answer, case study, and high-mark questions."""
    docx_file = _create_synthetic_long_answer_docx(tmp_path / "long_answer_test.docx")
    parser = DocxWorksheetParser()
    parsed = parser.parse(docx_file)

    assert len(parsed.questions) == 3
    for q in parsed.questions:
        assert q.question_type == QuestionType.LONG_ANSWER

    assert parsed.questions[0].marks == 10.0
    assert parsed.questions[1].marks == 16.0


def test_synthetic_mixed_worksheet(tmp_path: Path):
    """Verify handling of mixed worksheets across multiple sections."""
    docx_file = _create_synthetic_mixed_docx(tmp_path / "mixed_test.docx")
    parser = DocxWorksheetParser()
    parsed = parser.parse(docx_file)

    assert parsed.unit == 2
    assert parsed.session == 10
    assert parsed.slo == 2
    assert parsed.course_code == "21CSC303J"

    assert len(parsed.sections) == 3
    assert len(parsed.questions) == 4

    # Verify sequential ordering and distinct types
    q1, q2, q3, q4 = parsed.questions
    assert q1.question_number == "1"
    assert q1.question_type == QuestionType.MCQ
    assert q1.section == "Part A - Objective Section"

    assert q2.question_number == "2"
    assert q2.question_type == QuestionType.ONE_WORD

    assert q3.question_number == "3"
    assert q3.question_type == QuestionType.SHORT_ANSWER
    assert q3.marks == 2.0
    assert q3.section == "Part B - Conceptual Questions"

    assert q4.question_number == "4"
    assert q4.question_type == QuestionType.LONG_ANSWER
    assert q4.marks == 10.0
    assert q4.section == "Part C - Detailed Problems"


def test_synthetic_table_based_worksheet(tmp_path: Path):
    """Verify extracting questions laid out inside a table."""
    docx_file = _create_synthetic_table_docx(tmp_path / "table_test.docx")
    parser = DocxWorksheetParser()
    parsed = parser.parse(docx_file)

    assert len(parsed.questions) == 3

    q1, q2, q3 = parsed.questions
    assert q1.question_number == "1"
    assert "configuration management" in q1.question_text
    assert q1.marks == 2.0
    assert q1.question_type == QuestionType.SHORT_ANSWER

    assert q2.question_number == "2"
    assert "Spiral Model" in q2.question_text
    assert q2.marks == 10.0
    assert q2.question_type == QuestionType.LONG_ANSWER

    assert q3.question_number == "3"
    assert q3.question_type == QuestionType.MCQ
    assert len(q3.options) == 4


def test_docx_safe_round_trip(tmp_path: Path):
    """Verify parsing is completely non-destructive and order-deterministic."""
    original_file = _create_synthetic_mixed_docx(tmp_path / "original.docx")
    copied_file = tmp_path / "copy.docx"
    shutil.copyfile(original_file, copied_file)

    # Check hash before parsing
    hash_before = hashlib.sha256(original_file.read_bytes()).hexdigest()

    parser = DocxWorksheetParser()
    parsed_original = parser.parse(original_file)

    # Check hash after parsing
    hash_after = hashlib.sha256(original_file.read_bytes()).hexdigest()
    assert hash_before == hash_after, "Parsing must never modify the original file on disk!"

    # Parse copied file and assert complete structural equality
    parsed_copy = parser.parse(copied_file)
    assert len(parsed_original.questions) == len(parsed_copy.questions)
    for q_orig, q_copy in zip(parsed_original.questions, parsed_copy.questions):
        assert q_orig.question_number == q_copy.question_number
        assert q_orig.question_type == q_copy.question_type
        assert q_orig.question_text == q_copy.question_text
        assert q_orig.marks == q_copy.marks
        assert q_orig.source_order == q_copy.source_order


def test_real_1011_docx_if_present():
    """Inspect and verify the real SRM 1011.docx worksheet downloaded in Milestone 3."""
    real_path = Path("artifacts/real_1011.docx")
    if not real_path.exists():
        pytest.skip("Real 1011.docx artifact not found, skipping real file inspection test.")

    parser = DocxWorksheetParser()
    parsed = parser.parse(real_path)

    assert parsed.file_format == "docx"
    assert parsed.unit == 1
    assert parsed.session == 1
    assert parsed.slo == 1
    assert len(parsed.questions) == 2

    # Check Question 1 (Role-Play Simulation)
    q1 = parsed.questions[0]
    assert q1.question_number == "1"
    assert "Role-Play Simulation" in q1.question_text
    assert q1.question_type == QuestionType.LONG_ANSWER
    assert q1.context_or_activity is not None
    assert "Activity:" in q1.context_or_activity
    assert "Outcome:" in q1.context_or_activity

    # Check Question 2 (Design Thinking Workshop)
    q2 = parsed.questions[1]
    assert q2.question_number == "2"
    assert "Design Thinking Workshop" in q2.question_text
    assert q2.question_type == QuestionType.LONG_ANSWER
    assert q2.context_or_activity is not None


def test_pdf_parser_abstraction(tmp_path: Path):
    """Verify PDF worksheet parser extracts text, page locations, and questions."""
    pdf_path = tmp_path / "dummy.pdf"
    pdf_path.write_bytes(b"%PDF-1.4 mock")

    mock_page1 = MagicMock()
    mock_page1.extract_text.return_value = (
        "Unit I - Session 1 - SLO 1\n"
        "Part A - Multiple Choice\n"
        "1. Which language is typed dynamically?\n"
        "A. Python\n"
        "B. Java\n"
        "C. C++\n"
        "D. Rust\n"
    )

    mock_page2 = MagicMock()
    mock_page2.extract_text.return_value = (
        "Part B - Short Questions\n"
        "2. Define an operating system. [2 Marks]\n"
    )

    mock_reader = MagicMock()
    mock_reader.pages = [mock_page1, mock_page2]

    with patch("packages.worksheets.pdf_parser.PdfReader", return_value=mock_reader):
        parser = PdfWorksheetParser()
        parsed = parser.parse(pdf_path)

        assert parsed.file_format == "pdf"
        assert parsed.unit == 1
        assert parsed.session == 1
        assert parsed.slo == 1
        assert len(parsed.questions) == 2

        q1 = parsed.questions[0]
        assert q1.question_number == "1"
        assert q1.question_type == QuestionType.MCQ
        assert len(q1.options) == 4
        assert q1.source_location["page"] == 1

        q2 = parsed.questions[1]
        assert q2.question_number == "2"
        assert q2.question_type == QuestionType.SHORT_ANSWER
        assert q2.marks == 2.0
        assert q2.source_location["page"] == 2


def test_worksheet_processor_adapter(tmp_path: Path):
    """Verify DefaultWorksheetProcessor wraps parser and generates solutions."""
    docx_file = _create_synthetic_short_answer_docx(tmp_path / "processor_test.docx")
    processor = DefaultWorksheetProcessor()

    parsed_data = processor.parse_worksheet(docx_file)
    assert isinstance(parsed_data, dict)
    assert "questions" in parsed_data
    assert len(parsed_data["questions"]) == 4

    completed = processor.process_solutions(parsed_data)
    assert completed.exists()
    assert completed != docx_file
    assert completed.name == "completed_processor_test.docx"


def test_unified_worksheet_parser_dispatcher(tmp_path: Path):
    """Verify WorksheetParser dispatches to proper format engine and rejects invalid."""
    docx_file = _create_synthetic_short_answer_docx(tmp_path / "dispatch_test.docx")
    parser = WorksheetParser()

    parsed_docx = parser.parse(docx_file)
    assert parsed_docx.file_format == "docx"

    invalid_file = tmp_path / "test.txt"
    invalid_file.write_text("dummy")
    with pytest.raises(ValueError) as exc:
        parser.parse(invalid_file)
    assert "unsupported" in str(exc.value).lower()
