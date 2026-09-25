"""Tests for generic code-intent detection, language inference, and code formatting in worksheets."""

import pytest
from pathlib import Path
import docx

from packages.worksheets.models import (
    ParsedQuestion,
    ParsedWorksheet,
    QuestionType,
    ResponseMode,
    AnswerTargetSpec,
)
from packages.worksheets.answer_models import (
    GeneratedAnswer,
    WorksheetAnswers,
    AnswerStatus,
)
from packages.worksheets.classifier import CodeIntentDetector, QuestionClassifier
from packages.worksheets.answer_engine import RuleBasedAnswerEngine, LLMAnswerEngine
from packages.worksheets.answer_target import AnswerTargetResolver, TargetWriter
from packages.worksheets.filler import DocxWorksheetFiller


# ==============================================================================
# 1. SEMANTIC CODE-INTENT DETECTOR & LANGUAGE INFERENCE TESTS
# ==============================================================================

class TestCodeIntentDetector:
    """Verifies generic semantic detection of code intent, response mode, and language."""

    def test_q1_vehicle_hierarchy_detection(self):
        """'Design a class hierarchy: Vehicle → Car → ElectricCar.' -> CODE, java (via context)."""
        prompt = "Design a class hierarchy: Vehicle → Car → ElectricCar."
        mode, lang = CodeIntentDetector.detect(prompt, context={"course_code": "21CSC203P"})
        assert mode == ResponseMode.CODE
        assert lang == "java"

    def test_q2_employee_manager_detection(self):
        """'Create a superclass Employee and subclass Manager.' -> CODE, java (via context)."""
        prompt = "Create a superclass Employee and subclass Manager."
        mode, lang = CodeIntentDetector.detect(prompt, context={"course_code": "21CSC203P"})
        assert mode == ResponseMode.CODE
        assert lang == "java"

    def test_q3_override_and_explain_super_detection(self):
        """'Override a method in Java and explain the use of super.' -> CODE_AND_EXPLANATION, java."""
        prompt = "Override a method in Java and explain the use of super."
        mode, lang = CodeIntentDetector.detect(prompt)
        assert mode == ResponseMode.CODE_AND_EXPLANATION
        assert lang == "java"

    def test_c_programming_language_detection(self):
        """Explicit C language question."""
        prompt = "Write a C program to reverse an array using pointers."
        mode, lang = CodeIntentDetector.detect(prompt)
        assert mode == ResponseMode.CODE
        assert lang == "c"

    def test_python_programming_language_detection(self):
        """Explicit Python language question."""
        prompt = "Implement a binary search function in Python."
        mode, lang = CodeIntentDetector.detect(prompt)
        assert mode == ResponseMode.CODE
        assert lang == "python"

    def test_sql_query_detection(self):
        """Explicit SQL query question."""
        prompt = "Write an SQL query to find employees whose salary > 50000."
        mode, lang = CodeIntentDetector.detect(prompt)
        assert mode == ResponseMode.CODE
        assert lang == "sql"

    def test_pseudocode_detection(self):
        """Pseudocode prompt -> PSEUDOCODE."""
        prompt = "Write pseudocode for QuickSort."
        mode, lang = CodeIntentDetector.detect(prompt)
        assert mode == ResponseMode.PSEUDOCODE

    def test_algorithm_detection(self):
        """Algorithm prompt -> ALGORITHM."""
        prompt = "Write an algorithm to find the shortest path in a graph."
        mode, lang = CodeIntentDetector.detect(prompt)
        assert mode == ResponseMode.ALGORITHM

    def test_output_tracing_detection(self):
        """Output trace prompt -> OUTPUT_TRACE."""
        prompt = "Find the output of the following Java snippet: int x = 5; System.out.println(x++);"
        mode, lang = CodeIntentDetector.detect(prompt)
        assert mode == ResponseMode.OUTPUT_TRACE
        assert lang == "java"

    @pytest.mark.parametrize(
        "theory_prompt",
        [
            "Explain inheritance in Java with real-world examples.",
            "What is polymorphism? Describe runtime polymorphism.",
            "Compare abstract class and interface in Java.",
            "Define encapsulation and data hiding.",
            "Differentiate between method overloading and overriding.",
            "Discuss the primary features of Object-Oriented Programming.",
            "Write a short note on garbage collection in Java.",
            "Briefly explain the role of the JVM.",
        ],
    )
    def test_negative_pure_theory_prompts_remain_text(self, theory_prompt: str):
        """Theory/conceptual questions must NOT be classified as CODE."""
        mode, lang = CodeIntentDetector.detect(theory_prompt, context={"course_code": "21CSC203P"})
        assert mode == ResponseMode.TEXT, f"Failed for prompt: {theory_prompt}"


# ==============================================================================
# 2. QUESTION CLASSIFIER INTEGRATION & SEPARATION OF ANSWER_TYPE FROM RESPONSE_MODE
# ==============================================================================

class TestQuestionClassifierIntegration:
    """Verifies that QuestionClassifier separates question_type from response_mode."""

    def test_classify_preserves_short_answer_with_code_response_mode(self):
        """A 2-mark implementation question has question_type=SHORT_ANSWER and response_mode=CODE."""
        q = ParsedQuestion(
            question_id="q1",
            question_number="1",
            question_text="Design a class hierarchy: Vehicle → Car → ElectricCar. [2 marks]",
        )
        qtype = QuestionClassifier.classify(q, context={"course_code": "21CSC203P"})

        # question_type is preserved as SHORT_ANSWER based on 2 marks
        assert qtype == QuestionType.SHORT_ANSWER
        assert q.question_type == QuestionType.SHORT_ANSWER
        # response_mode is correctly detected as CODE
        assert q.response_mode == ResponseMode.CODE
        assert q.language == "java"

    def test_classify_preserves_long_answer_with_code_and_explanation(self):
        """An 8-mark question has question_type=LONG_ANSWER and response_mode=CODE_AND_EXPLANATION."""
        q = ParsedQuestion(
            question_id="q2",
            question_number="2",
            question_text="Override a method in Java and explain the use of super. [8 marks]",
        )
        qtype = QuestionClassifier.classify(q, context={"course_code": "21CSC203P"})

        assert qtype == QuestionType.LONG_ANSWER
        assert q.question_type == QuestionType.LONG_ANSWER
        assert q.response_mode == ResponseMode.CODE_AND_EXPLANATION
        assert q.language == "java"

    def test_classify_table_activity_response_mode(self):
        """Table cells get TABLE_VALUE response mode and compact available space."""
        target = AnswerTargetSpec(
            target_id="tbl_1_c1",
            target_type="table_cell",
            expected_length="phrase",
        )
        q = ParsedQuestion(
            question_id="q3",
            question_text="Fill the comparison table.",
            question_type=QuestionType.TABLE_CELL,
            targets=[target],
        )
        qtype = QuestionClassifier.classify(q)
        assert q.response_mode == ResponseMode.TABLE_VALUE
        assert q.available_space == "compact"


# ==============================================================================
# 3. RULE-BASED ENGINE CODE GENERATION FOR 21CSC203P QUESTIONS
# ==============================================================================

class TestRuleBasedEngineCodeGeneration:
    """Verifies that the offline/test rule engine generates valid Java code for the target questions."""

    @pytest.mark.asyncio
    async def test_rule_engine_q1_vehicle_hierarchy(self):
        engine = RuleBasedAnswerEngine()
        q = ParsedQuestion(
            question_id="q1",
            question_text="Design a class hierarchy: Vehicle → Car → ElectricCar.",
            question_type=QuestionType.SHORT_ANSWER,
            response_mode=ResponseMode.CODE,
            language="java",
        )
        ans = await engine.answer_question(q)
        assert ans.status == AnswerStatus.SUCCESS
        assert ans.response_mode == ResponseMode.CODE
        assert "class Vehicle" in ans.answer_text
        assert "class Car extends Vehicle" in ans.answer_text
        assert "class ElectricCar extends Car" in ans.answer_text

    @pytest.mark.asyncio
    async def test_rule_engine_q2_employee_manager(self):
        engine = RuleBasedAnswerEngine()
        q = ParsedQuestion(
            question_id="q2",
            question_text="Create a superclass Employee and subclass Manager.",
            question_type=QuestionType.SHORT_ANSWER,
            response_mode=ResponseMode.CODE,
            language="java",
        )
        ans = await engine.answer_question(q)
        assert ans.status == AnswerStatus.SUCCESS
        assert ans.response_mode == ResponseMode.CODE
        assert "class Employee" in ans.answer_text
        assert "class Manager extends Employee" in ans.answer_text

    @pytest.mark.asyncio
    async def test_rule_engine_q3_override_super(self):
        engine = RuleBasedAnswerEngine()
        q = ParsedQuestion(
            question_id="q3",
            question_text="Override a method in Java and explain the use of super.",
            question_type=QuestionType.SHORT_ANSWER,
            response_mode=ResponseMode.CODE_AND_EXPLANATION,
            language="java",
        )
        ans = await engine.answer_question(q)
        assert ans.status == AnswerStatus.SUCCESS
        assert ans.response_mode == ResponseMode.CODE_AND_EXPLANATION
        assert "@Override" in ans.answer_text
        assert "super." in ans.answer_text
        assert "Explanation:" in ans.answer_text


# ==============================================================================
# 4. LLM ENGINE PAYLOAD AND PARSING INTEGRATION
# ==============================================================================

class TestLLMEnginePayloadAndParsing:
    """Verifies response_mode serialization in LLM payloads and parsing back into GeneratedAnswer."""

    def test_build_questions_payload_includes_response_mode_and_language(self):
        engine = LLMAnswerEngine(allow_fallback_when_unconfigured=True)
        q1 = ParsedQuestion(
            question_id="q1",
            question_number="1",
            question_text="Design a class hierarchy: Vehicle → Car → ElectricCar.",
            question_type=QuestionType.SHORT_ANSWER,
            response_mode=ResponseMode.CODE,
            language="java",
            available_space="complete",
        )
        ws = ParsedWorksheet(
            filename="sample.docx",
            file_format="docx",
            course_code="21CSC203P",
            questions=[q1],
        )
        payload = engine._build_questions_payload(ws)
        assert len(payload) == 1
        item = payload[0]
        assert item["question_id"] == "q1"
        assert item["question_type"] == "SHORT_ANSWER"
        assert item["response_mode"] == "CODE"
        assert item["language"] == "java"
        assert item["available_space"] == "complete"

    def test_map_and_validate_answers_parses_response_mode_and_language(self):
        engine = LLMAnswerEngine(allow_fallback_when_unconfigured=True)
        q1 = ParsedQuestion(
            question_id="q1",
            question_number="1",
            question_text="Design a class hierarchy: Vehicle → Car → ElectricCar.",
            question_type=QuestionType.SHORT_ANSWER,
            response_mode=ResponseMode.CODE,
            language="java",
        )
        ws = ParsedWorksheet(
            filename="sample.docx",
            file_format="docx",
            course_code="21CSC203P",
            questions=[q1],
        )
        raw_ai_answers = [
            {
                "question_id": "q1",
                "response_mode": "CODE",
                "language": "java",
                "answer_text": "class Vehicle {}\nclass Car extends Vehicle {}",
                "confidence": 0.96,
            }
        ]
        result = engine._map_and_validate_answers(raw_ai_answers, ws)
        assert len(result.answers) == 1
        ans = result.answers[0]
        assert ans.response_mode == ResponseMode.CODE
        assert ans.language == "java"
        assert ans.answer == "class Vehicle {}\nclass Car extends Vehicle {}"
        assert ans.answer_text == "class Vehicle {}\nclass Car extends Vehicle {}"


# ==============================================================================
# 5. DOCX FILLER CODE FORMATTING & INDENTATION PRESERVATION
# ==============================================================================

class TestDocxFillerCodeFormatting:
    """Verifies that the DOCX filler preserves leading indentation and applies Consolas font."""

    def test_writer_preserves_indentation_and_applies_consolas(self, tmp_path: Path):
        doc = docx.Document()
        p = doc.add_paragraph("Answer: ")

        writer = TargetWriter()
        q = ParsedQuestion(
            question_id="q1",
            question_number="1",
            question_text="Create a superclass Employee and subclass Manager.",
            question_type=QuestionType.SHORT_ANSWER,
            response_mode=ResponseMode.CODE,
            language="java",
        )
        ans = GeneratedAnswer(
            question_id="q1",
            question_number="1",
            question_type=QuestionType.SHORT_ANSWER,
            response_mode=ResponseMode.CODE,
            language="java",
            answer_text=(
                "class Employee {\n"
                "    protected String name;\n"
                "    public Employee(String name) {\n"
                "        this.name = name;\n"
                "    }\n"
                "}"
            ),
        )

        target = writer.resolve if hasattr(writer, "resolve") else None
        # Write into placeholder
        writer._write_paragraph_placeholder(p, ans.answer_text.splitlines(), placeholder_text="Answer: ", is_code=True)

        # Check paragraphs
        paragraphs = doc.paragraphs
        assert len(paragraphs) >= 5

        # Check that indentation is preserved in paragraphs after the first line
        # Line 1: '    protected String name;' -> should have 4 leading spaces
        second_p_text = paragraphs[1].text
        assert second_p_text.startswith("    protected")

        # Check that Consolas font is applied to runs
        for p_elem in paragraphs[1:]:
            for r in p_elem.runs:
                assert r.font.name == "Consolas"
                assert r.font.size.pt == 9.0

    def test_writer_code_and_explanation_splits_fonts(self, tmp_path: Path):
        """Verifies code is rendered in Consolas while explanation is rendered in standard font."""
        doc = docx.Document()
        p = doc.add_paragraph()

        writer = TargetWriter()
        q = ParsedQuestion(
            question_id="q3",
            question_text="Override a method in Java and explain the use of super.",
            question_type=QuestionType.SHORT_ANSWER,
            response_mode=ResponseMode.CODE_AND_EXPLANATION,
            language="java",
        )
        code_and_expl = (
            "class Animal {\n"
            "    public void sound() {}\n"
            "}\n"
            "Explanation:\n"
            "Method overriding allows a subclass to provide a specific implementation."
        )
        ans = GeneratedAnswer(
            question_id="q3",
            question_type=QuestionType.SHORT_ANSWER,
            response_mode=ResponseMode.CODE_AND_EXPLANATION,
            language="java",
            answer_text=code_and_expl,
        )

        writer._write_paragraph_empty(p, code_and_expl.splitlines(), q, is_code=True)

        paragraphs = doc.paragraphs
        assert len(paragraphs) == 5

        # Paragraphs 0, 1, 2 are code -> Consolas
        for p_elem in paragraphs[:3]:
            for r in p_elem.runs:
                assert r.font.name == "Consolas"

        # Paragraphs 3 and 4 are Explanation -> standard font (font.name is not Consolas, font.size 10.0)
        for p_elem in paragraphs[3:]:
            for r in p_elem.runs:
                assert r.font.name != "Consolas"
                assert r.font.size.pt == 10.0
