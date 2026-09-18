"""Comprehensive regression and unit tests for generic worksheet answer-target filling.

Verifies:
1. Real 1052.docx (SLO2) structural pattern regression:
   - All 4 questions receive answers.
   - Each answer appears in the intended answer box/location.
   - No answer is written into question text.
   - No answer is duplicated.
   - No answer target is silently skipped.
   - SLO2 worksheet filename remains 1052.docx.
   - Original 1052.docx SHA256 remains completely unchanged.
   - Completed file opens successfully and XML contains answers at exact locations.
2. Paragraph answer field (explicit placeholder).
3. Table-cell answer field (dedicated answer column).
4. Bordered/blank answer cell (row-alternating table pattern).
5. Text-box answer field (w:txbxContent in DrawingML/VML shape).
6. Content-control answer field (w:sdt / w:sdtContent).
7. Structured error handling (raises WorksheetFillingError when target cannot be resolved).
"""

import hashlib
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
import pytest
import docx
from docx.oxml import parse_xml
from docx.shared import Pt, RGBColor

from packages.worksheets.answer_models import GeneratedAnswer, WorksheetAnswers
from packages.worksheets.answer_target import (
    ANSWER_COLOR_RGB,
    AnswerTargetResolver,
    AnswerTargetType,
    TargetWriter,
)
from packages.worksheets.docx_parser import DocxWorksheetParser
from packages.worksheets.exceptions import WorksheetFillingError
from packages.worksheets.filler import DocxWorksheetFiller
from packages.worksheets.models import ParsedQuestion, QuestionType


# ---------------------------------------------------------------------------
# Test 1: Real 1052.docx SLO2 Structural Pattern Regression
# ---------------------------------------------------------------------------

def test_slo2_1052_structural_pattern_regression(tmp_path: Path):
    """Verify end-to-end target resolution and filling on real 1052.docx SLO2 worksheet."""
    source_1052 = Path("artifacts/1052.docx")
    assert source_1052.exists(), "artifacts/1052.docx must exist for regression test"

    # Make a copy in tmp_path to test immutability guarantee in isolation
    test_1052 = tmp_path / "1052.docx"
    test_1052.write_bytes(source_1052.read_bytes())
    sha_before = hashlib.sha256(test_1052.read_bytes()).hexdigest()

    parser = DocxWorksheetParser()
    ws = parser.parse(test_1052)

    # 1. Verify 4 questions parsed and target metadata captured
    assert len(ws.questions) == 4, f"Expected 4 questions, got {len(ws.questions)}"
    for q in ws.questions:
        assert q.question_text, f"Question {q.question_id} has empty text"
        # Verify no answer was swallowed into the question text
        assert not q.question_text.strip().endswith("Answer:"), (
            f"Question {q.question_id} had placeholder 'Answer:' swallowed into question text!"
        )

    # Verify placeholder locations in 1052.docx
    assert ws.questions[0].source_location.get("answer_placeholder_paragraph_index") == 3
    assert ws.questions[1].source_location.get("answer_placeholder_paragraph_index") == 5
    assert ws.questions[2].source_location.get("answer_placeholder_paragraph_index") == 7
    assert ws.questions[3].source_location.get("answer_placeholder_paragraph_index") == 9

    # 2. Build answers
    answers = WorksheetAnswers(
        worksheet_title=ws.title,
        worksheet_filename="1052.docx",
        answers=[
            GeneratedAnswer(
                question_id="q_1_1",
                question_number="1",
                question_type=QuestionType.MCQ,
                selected_option="A",
                answer_text="Model-View-Controller (MVC)",
                confidence=1.0,
            ),
            GeneratedAnswer(
                question_id="q_2_2",
                question_number="2",
                question_type=QuestionType.ONE_WORD,
                answer_text="Representational State Transfer",
                confidence=1.0,
            ),
            GeneratedAnswer(
                question_id="q_3_3",
                question_number="3",
                question_type=QuestionType.SHORT_ANSWER,
                answer_text="Monolithic architecture bundles all components into a single unit, whereas microservices partition functionality into modular services.",
                confidence=1.0,
            ),
            GeneratedAnswer(
                question_id="q_4_4",
                question_number="4",
                question_type=QuestionType.LONG_ANSWER,
                answer_text="Scalable Campus Notification System Architecture:\n1. Ingestion: Gateway receives alerts.\n2. Broker: Distributed queue buffers messages.\n3. Dispatchers: Multi-channel workers send notifications.",
                confidence=1.0,
            ),
        ],
    )

    filler = DocxWorksheetFiller()
    out_dir = tmp_path / "output"
    completed_file = filler.fill(test_1052, ws, answers, output_dir=out_dir)

    # 3. Verify original file SHA256 was 100% unchanged
    sha_after = hashlib.sha256(test_1052.read_bytes()).hexdigest()
    assert sha_before == sha_after, "Original 1052.docx SHA256 changed during filling!"

    # 4. Verify completed file properties
    assert completed_file.exists()
    assert completed_file.name == "completed_1052.docx"

    # Open completed document
    comp_doc = docx.Document(str(completed_file))
    paragraphs = comp_doc.paragraphs

    # Verify answers are in their intended paragraph locations:
    # P[2] = Q1, P[3] = Answer 1
    assert "Which architectural pattern" in paragraphs[2].text
    assert paragraphs[3].text.strip() == "Answer: A. Model-View-Controller (MVC)"

    # P[4] = Q2, P[5] = Answer 2
    assert "Expand the architectural acronym REST" in paragraphs[4].text
    assert paragraphs[5].text.strip() == "Answer: Representational State Transfer"

    # P[6] = Q3, P[7] = Answer 3
    assert "Differentiate between Monolithic Architecture" in paragraphs[6].text
    assert "Answer: Monolithic architecture bundles" in paragraphs[7].text

    # P[8] = Q4, P[9..12] = Answer 4
    assert "Case Study Workshop" in paragraphs[8].text
    assert "Answer: Scalable Campus Notification System" in paragraphs[9].text
    assert "1. Ingestion: Gateway receives alerts." in paragraphs[10].text
    assert "2. Broker: Distributed queue buffers messages." in paragraphs[11].text
    assert "3. Dispatchers: Multi-channel workers send notifications." in paragraphs[12].text

    # Verify no duplicate empty "Answer:" paragraphs exist
    empty_answer_paras = [p.text for p in paragraphs if p.text.strip() in ("Answer:", "Answer :")]
    assert len(empty_answer_paras) == 0, f"Found duplicate empty answer paragraphs: {empty_answer_paras}"

    # Verify physical presence in XML
    with zipfile.ZipFile(str(completed_file)) as z:
        xml_bytes = z.read("word/document.xml")
    xml_str = xml_bytes.decode("utf-8")
    assert "Model-View-Controller (MVC)" in xml_str
    assert "Representational State Transfer" in xml_str
    assert "1F4E79" in xml_str, "Answer styling color (1F4E79) not found in document XML!"


# ---------------------------------------------------------------------------
# Test 2: Paragraph Answer Field Filling
# ---------------------------------------------------------------------------

def test_paragraph_answer_field_filling(tmp_path: Path):
    """Verify paragraph with explicit 'Solution:' prompt is filled cleanly without duplicate paragraphs."""
    doc_path = tmp_path / "para_test.docx"
    doc = docx.Document()
    doc.add_heading("Unit 1 - Session 1 - SLO 1", level=1)
    doc.add_paragraph("1. Explain the Single Responsibility Principle.")
    doc.add_paragraph("Solution: ")
    doc.save(str(doc_path))

    parser = DocxWorksheetParser()
    ws = parser.parse(doc_path)
    assert len(ws.questions) == 1

    answers = WorksheetAnswers(
        worksheet_title="Test",
        worksheet_filename="para_test.docx",
        answers=[
            GeneratedAnswer(
                question_id=ws.questions[0].question_id,
                question_number="1",
                question_type=QuestionType.SHORT_ANSWER,
                answer_text="A class should have only one reason to change.",
                confidence=1.0,
            )
        ],
    )

    filler = DocxWorksheetFiller()
    comp_path = filler.fill(doc_path, ws, answers, output_dir=tmp_path)

    comp_doc = docx.Document(str(comp_path))
    # Expect 3 paragraphs: Heading, Question, Answer (no duplicate empty placeholder)
    assert len(comp_doc.paragraphs) == 3
    assert comp_doc.paragraphs[1].text.startswith("1. Explain")
    assert comp_doc.paragraphs[2].text == "Answer: A class should have only one reason to change."


# ---------------------------------------------------------------------------
# Test 3: Table-Cell Answer Field Filling (Dedicated Answer Column)
# ---------------------------------------------------------------------------

def test_table_cell_answer_column_filling(tmp_path: Path):
    """Verify table with dedicated 'Answer' column fills answers directly in the answer cell."""
    doc_path = tmp_path / "table_col_test.docx"
    doc = docx.Document()
    doc.add_heading("Table Questions", level=1)

    tbl = doc.add_table(rows=3, cols=3)
    # Header
    tbl.rows[0].cells[0].text = "Q.No"
    tbl.rows[0].cells[1].text = "Question Description"
    tbl.rows[0].cells[2].text = "Answer"

    # Row 1
    tbl.rows[1].cells[0].text = "1"
    tbl.rows[1].cells[1].text = "What is encapsulation?"
    tbl.rows[1].cells[2].text = ""  # Empty answer cell

    # Row 2
    tbl.rows[2].cells[0].text = "2"
    tbl.rows[2].cells[1].text = "What is polymorphism?"
    tbl.rows[2].cells[2].text = ""  # Empty answer cell

    doc.save(str(doc_path))

    parser = DocxWorksheetParser()
    ws = parser.parse(doc_path)
    assert len(ws.questions) == 2

    answers = WorksheetAnswers(
        worksheet_title="Table Test",
        worksheet_filename="table_col_test.docx",
        answers=[
            GeneratedAnswer(
                question_id=ws.questions[0].question_id,
                question_number="1",
                question_type=QuestionType.SHORT_ANSWER,
                answer_text="Encapsulation bundles data and methods.",
                confidence=1.0,
            ),
            GeneratedAnswer(
                question_id=ws.questions[1].question_id,
                question_number="2",
                question_type=QuestionType.SHORT_ANSWER,
                answer_text="Polymorphism allows multiple forms.",
                confidence=1.0,
            ),
        ],
    )

    filler = DocxWorksheetFiller()
    comp_path = filler.fill(doc_path, ws, answers, output_dir=tmp_path)

    comp_doc = docx.Document(str(comp_path))
    table = comp_doc.tables[0]

    # Verify answers were written into Column 2, leaving question in Column 1 intact
    assert table.rows[1].cells[1].text == "What is encapsulation?"
    assert "Encapsulation bundles data and methods." in table.rows[1].cells[2].text

    assert table.rows[2].cells[1].text == "What is polymorphism?"
    assert "Polymorphism allows multiple forms." in table.rows[2].cells[2].text


# ---------------------------------------------------------------------------
# Test 4: Bordered/Blank Answer Cell (Row-Alternating Table Pattern)
# ---------------------------------------------------------------------------

def test_bordered_blank_cell_row_alternating_filling(tmp_path: Path):
    """Verify table where row i is Question and row i+1 is a bordered blank answer box."""
    doc_path = tmp_path / "alternating_table_test.docx"
    doc = docx.Document()
    doc.add_heading("Exam Questions", level=1)

    tbl = doc.add_table(rows=2, cols=1)
    # Row 0: Question
    tbl.rows[0].cells[0].text = "1. Describe the three-tier software architecture."

    # Row 1: Bordered answer box with placeholder
    c1 = tbl.rows[1].cells[0]
    c1.text = "Answer: "
    # Add border xml
    tcPr = c1._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        '<w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        '</w:tcBorders>'
    )
    tcPr.append(tcBorders)

    doc.save(str(doc_path))

    # Construct ParsedQuestion representing row 0
    pq = ParsedQuestion(
        question_id="tbl_q_1",
        question_number="1",
        question_text="Describe the three-tier software architecture.",
        question_type=QuestionType.SHORT_ANSWER,
        source_location={"table_index": 0, "row_index": 0, "col_index": 0},
    )
    ws = docx_worksheet = docx.Document(str(doc_path))

    resolver = AnswerTargetResolver()
    target = resolver.resolve(ws, pq, [pq])

    assert target.target_type == AnswerTargetType.TABLE_CELL_ROW
    assert target.row_index == 1, "Expected target row to be row 1 (alternating answer cell)"

    writer = TargetWriter()
    ans = GeneratedAnswer(
        question_id="tbl_q_1",
        question_number="1",
        question_type=QuestionType.SHORT_ANSWER,
        answer_text="Presentation, Application Logic, and Data Storage tiers.",
        confidence=1.0,
    )
    writer.write(ws, target, pq, ans)

    # Verify written content
    assert ws.tables[0].rows[0].cells[0].text == "1. Describe the three-tier software architecture."
    assert "Presentation, Application Logic, and Data Storage" in ws.tables[0].rows[1].cells[0].text


# ---------------------------------------------------------------------------
# Test 5: Text-Box Answer Field Filling (w:txbxContent)
# ---------------------------------------------------------------------------

def test_text_box_answer_field_filling(tmp_path: Path):
    """Verify writing into w:txbxContent inside a drawing/shape answer box."""
    doc_path = tmp_path / "textbox_test.docx"
    doc = docx.Document()
    doc.add_paragraph("1. Explain the difference between synchronous and asynchronous calls.")

    # Create paragraph containing a drawing text box
    tb_p = doc.add_paragraph()
    xml_str = (
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<w:drawing>'
        '<wp:inline>'
        '<wps:wsp>'
        '<wps:txbx>'
        '<w:txbxContent>'
        '<w:p><w:r><w:t>Write answer here...</w:t></w:r></w:p>'
        '</w:txbxContent>'
        '</wps:txbx>'
        '</wps:wsp>'
        '</wp:inline>'
        '</w:drawing>'
        '</w:r>'
    )
    tb_r = parse_xml(xml_str)
    tb_p._p.append(tb_r)

    doc.save(str(doc_path))

    pq = ParsedQuestion(
        question_id="tb_q_1",
        question_number="1",
        question_text="Explain the difference between synchronous and asynchronous calls.",
        question_type=QuestionType.SHORT_ANSWER,
        source_location={"paragraph_index": 0},
    )

    doc_loaded = docx.Document(str(doc_path))
    resolver = AnswerTargetResolver()
    target = resolver.resolve(doc_loaded, pq, [pq])

    assert target.target_type == AnswerTargetType.TEXT_BOX
    assert target.txbx_element is not None

    writer = TargetWriter()
    ans = GeneratedAnswer(
        question_id="tb_q_1",
        question_number="1",
        question_type=QuestionType.SHORT_ANSWER,
        answer_text="Synchronous blocks until completion; asynchronous returns immediately.",
        confidence=1.0,
    )
    writer.write(doc_loaded, target, pq, ans)

    # Verify text box was physically updated in XML
    txbx_p = target.txbx_element.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p")[0]
    p_wrapper = docx.text.paragraph.Paragraph(txbx_p, doc_loaded)
    assert "Synchronous blocks until completion" in p_wrapper.text
    assert "Write answer here..." not in p_wrapper.text


# ---------------------------------------------------------------------------
# Test 6: Content-Control Answer Field Filling (w:sdt / w:sdtContent)
# ---------------------------------------------------------------------------

def test_content_control_answer_field_filling(tmp_path: Path):
    """Verify writing into w:sdtContent structured document tag."""
    doc_path = tmp_path / "sdt_test.docx"
    doc = docx.Document()
    doc.add_paragraph("1. Define recursion.")

    sdt_p = doc.add_paragraph()
    sdt_xml = (
        '<w:sdt xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:sdtPr><w:alias w:val="AnswerBox"/></w:sdtPr>'
        '<w:sdtContent>'
        '<w:p><w:r><w:t>Click or tap here to enter text.</w:t></w:r></w:p>'
        '</w:sdtContent>'
        '</w:sdt>'
    )
    sdt_elem = parse_xml(sdt_xml)
    sdt_p._p.append(sdt_elem)

    doc.save(str(doc_path))

    pq = ParsedQuestion(
        question_id="sdt_q_1",
        question_number="1",
        question_text="Define recursion.",
        question_type=QuestionType.SHORT_ANSWER,
        source_location={"paragraph_index": 0},
    )

    doc_loaded = docx.Document(str(doc_path))
    resolver = AnswerTargetResolver()
    target = resolver.resolve(doc_loaded, pq, [pq])

    assert target.target_type == AnswerTargetType.CONTENT_CONTROL
    assert target.sdt_element is not None

    writer = TargetWriter()
    ans = GeneratedAnswer(
        question_id="sdt_q_1",
        question_number="1",
        question_type=QuestionType.SHORT_ANSWER,
        answer_text="Recursion is a process where a function calls itself.",
        confidence=1.0,
    )
    writer.write(doc_loaded, target, pq, ans)

    # Check updated SDT text
    sdt_p_elem = target.sdt_element.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p")[0]
    p_wrap = docx.text.paragraph.Paragraph(sdt_p_elem, doc_loaded)
    assert "Recursion is a process where a function calls itself." in p_wrap.text


# ---------------------------------------------------------------------------
# Test 7: Error on Unresolvable Target
# ---------------------------------------------------------------------------

def test_error_raised_on_unresolvable_target(tmp_path: Path):
    """Verify WorksheetFillingError is raised when no target can be confidently resolved."""
    doc = docx.Document()
    doc.add_paragraph("Header only document")

    # Question points to nonexistent paragraph index and has no match
    corrupted_q = ParsedQuestion(
        question_id="corrupted_q_99",
        question_number="99",
        question_text="Completely missing question that does not exist anywhere.",
        question_type=QuestionType.SHORT_ANSWER,
        source_location={"paragraph_index": 999},
    )

    resolver = AnswerTargetResolver()
    with pytest.raises(WorksheetFillingError) as exc_info:
        resolver.resolve(doc, corrupted_q, [corrupted_q])

    assert "No writable target could be resolved" in str(exc_info.value)
