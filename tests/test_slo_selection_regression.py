"""Regression tests for SLO 1 vs SLO 2 selection and data flow.

Verifies:
1. Exact SLO 2 selection throughout the 10-stage pipeline:
   - Discovery response returns [SLO 1, SLO 2] (SLO 1 at index 0).
   - Selection preserves SLO 2 (worksheet_id 1052).
   - Worker targets 1052.docx, submits SLO=2, verifies SLO=2.
   - Result preserves session=105, slo=2, original_file="1052.docx".
2. Various worksheet_id format parsers:
   - "1052", "105_2", "105-2", "1052.docx", "Unit 1 Session 5 SLO 2".
3. Idempotency isolation between SLO 1 and SLO 2:
   - Verified status on SLO 1 (1051: 2) does not falsely mark SLO 2 (1052: 0) as already submitted.
4. HTTP Client verification candidate keys:
   - verify_submission correctly checks SLO 2 link and status, not SLO 1.
"""

from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch
import pytest
from docx import Document
from sqlalchemy.orm import Session

from apps.worker.tasks import _run_job_workflow, _run_submission_workflow
from packages.drive.client import GoogleDriveClient
from packages.drive.models import DriveFileMetadata
from packages.shared.models.job import Job, JobStatus
from packages.srm.http_client import SRMHttpClient
from packages.srm.models import (
    SRMCourse,
    SRMSessionStatus,
    SRMSubmissionResult,
    SRMWorksheetMetadata,
)
from packages.srm.orchestrator import SRMOrchestrator


def _setup_mock_orchestrator_for_session(session: int = 105):
    """Create mock SRMOrchestrator with both SLO 1 and SLO 2 available."""
    mock_orch = MagicMock(spec=SRMOrchestrator)
    mock_orch.transport_name = "http"
    mock_orch.connect = AsyncMock(return_value=True)
    mock_orch.authenticate = AsyncMock(return_value=True)
    mock_orch.capture_login_captcha = AsyncMock(return_value=None)
    mock_orch.close = AsyncMock(return_value=None)

    course = SRMCourse(
        course_code="21CSC303J",
        course_name="Software Engineering and Architecture",
        batch_id="B1",
        semester=3,
        department="CSE",
    )
    mock_orch.get_courses = AsyncMock(return_value=[course])
    mock_orch.get_courses_by_semester = AsyncMock(return_value=[course])

    ws_slo1 = SRMWorksheetMetadata(
        course_code="21CSC303J",
        session=session,
        slo=1,
        filename=f"{session}1.docx",
        format="docx",
        storage_path="data/coordinator/21CSC303J/slp",
        download_url=f"https://srm.portal/uploads/{session}1.docx",
        is_available=True,
        submission_status="NOT_SUBMITTED",
        title="Unit 1 Session 5 SLO 1",
    )
    ws_slo2 = SRMWorksheetMetadata(
        course_code="21CSC303J",
        session=session,
        slo=2,
        filename=f"{session}2.docx",
        format="docx",
        storage_path="data/coordinator/21CSC303J/slp",
        download_url=f"https://srm.portal/uploads/{session}2.docx",
        is_available=True,
        submission_status="NOT_SUBMITTED",
        title="Unit 1 Session 5 SLO 2",
    )

    # Note: index 0 is ws_slo1 (SLO 1)
    mock_orch.discover_worksheets = AsyncMock(return_value=[ws_slo1, ws_slo2])

    async def fake_get_file(course_code, session=1, slo=1, format_type="docx", filename=None):
        return f"https://srm.portal/uploads/{session}{slo}.docx"

    mock_orch.get_worksheet_file = AsyncMock(side_effect=fake_get_file)

    async def fake_download(file_url_or_id, destination_dir=None, filename=None):
        target = Path(destination_dir or ".") / (filename or "ws.docx")
        doc = Document()
        doc.add_heading(f"21CSC303J Session {session}", level=1)
        doc.add_paragraph("Explain the core architectural concepts.")
        doc.save(str(target))
        return target

    mock_orch.download_worksheet = AsyncMock(side_effect=fake_download)
    mock_orch.get_session_status = AsyncMock(
        return_value=SRMSessionStatus(
            session=session,
            practice_status={f"{session}1": 0, f"{session}2": 0},
            slo_links={},
        )
    )
    mock_orch.submit_worksheet_link = AsyncMock(
        return_value=SRMSubmissionResult(
            success=True,
            message="Submitted successfully",
            returned_link="https://docs.google.com/document/d/1A2B3C4D5E6F7G8H9I0J_mock/edit",
        )
    )
    mock_orch.verify_submission = AsyncMock(return_value=True)
    return mock_orch


def _setup_mock_drive():
    mock_drive = MagicMock(spec=GoogleDriveClient)
    mock_drive.upload_file = AsyncMock(
        return_value=DriveFileMetadata(
            file_id="1A2B3C4D5E6F7G8H9I0J_drive_file",
            filename="completed.docx",
            mime_type="application/vnd.google-apps.document",
            web_url="https://docs.google.com/document/d/1A2B3C4D5E6F7G8H9I0J_drive_file/edit",
            permission_status="VERIFIED_PUBLIC_READER",
            is_public=True,
        )
    )
    return mock_drive


@pytest.mark.asyncio
async def test_slo2_selection_workflow_end_to_end(db_session: Session):
    """Verify that selecting SLO 2 processes 1052.docx and never falls back to SLO 1 (1051.docx)."""
    db = db_session
    job = Job(
        user_id="RA2111003010001",
        course_id="21CSC303J",
        semester_id="3",
        worksheet_id="1052",  # SLO 2 explicitly selected
        transport_mode="http",
        status=JobStatus.PENDING,
    )
    db.add(job)
    db.commit()
    db.refresh(job)

    mock_orch = _setup_mock_orchestrator_for_session(105)
    mock_drive = _setup_mock_drive()

    creds = {
        "USER_ID": "RA2111003010001",
        "PASSWORD": "secretpassword",
        "requested_session": 105,
        "requested_slo": 2,
    }

    with patch("apps.worker.tasks.store_job_credentials"), \
         patch("apps.worker.tasks.get_job_credentials", return_value=creds), \
         patch("apps.worker.tasks.clear_job_credentials"):
        await _run_job_workflow(
            job_id=job.id,
            credentials=creds,
            db_session=db,
            orchestrator=mock_orch,
            drive_client=mock_drive,
        )

        db.refresh(job)
        assert job.status == JobStatus.AWAITING_USER_REVIEW

        await _run_submission_workflow(
            job_id=job.id,
            credentials=creds,
            db_session=db,
            orchestrator=mock_orch,
            drive_client=mock_drive,
        )

    db.refresh(job)
    assert job.status == JobStatus.COMPLETED
    assert job.result is not None
    assert job.result["session"] == 105
    assert job.result["slo"] == 2
    assert "1052" in job.result["original_file"]
    assert "1052" in job.result["completed_file"]

    # Verify orchestrator calls targeted SLO 2 specifically:
    # 1. download_worksheet should have downloaded 1052.docx
    mock_orch.download_worksheet.assert_called_once()
    dl_kwargs = mock_orch.download_worksheet.call_args[1]
    assert dl_kwargs["filename"] == "1052.docx"

    # 2. submit_worksheet_link should have passed session=105, slo=2
    mock_orch.submit_worksheet_link.assert_called_once()
    sub_kwargs = mock_orch.submit_worksheet_link.call_args[1]
    assert sub_kwargs["session"] == 105
    assert sub_kwargs["slo"] == 2

    # 3. verify_submission should have checked session=105, slo=2
    mock_orch.verify_submission.assert_called_once()
    ver_kwargs = mock_orch.verify_submission.call_args[1]
    assert ver_kwargs["session_or_worksheet_id"] == 105
    assert ver_kwargs["slo"] == 2


@pytest.mark.asyncio
async def test_slo1_selection_workflow_end_to_end(db_session: Session):
    """Verify that selecting SLO 1 processes 1051.docx."""
    db = db_session
    job = Job(
        user_id="RA2111003010001",
        course_id="21CSC303J",
        semester_id="3",
        worksheet_id="1051",  # SLO 1 explicitly selected
        transport_mode="http",
        status=JobStatus.PENDING,
    )
    db.add(job)
    db.commit()
    db.refresh(job)

    mock_orch = _setup_mock_orchestrator_for_session(105)
    mock_drive = _setup_mock_drive()

    creds = {
        "USER_ID": "RA2111003010001",
        "PASSWORD": "secretpassword",
        "requested_session": 105,
        "requested_slo": 1,
    }

    with patch("apps.worker.tasks.store_job_credentials"), \
         patch("apps.worker.tasks.get_job_credentials", return_value=creds), \
         patch("apps.worker.tasks.clear_job_credentials"):
        await _run_job_workflow(
            job_id=job.id,
            credentials=creds,
            db_session=db,
            orchestrator=mock_orch,
            drive_client=mock_drive,
        )

        db.refresh(job)
        assert job.status == JobStatus.AWAITING_USER_REVIEW

        await _run_submission_workflow(
            job_id=job.id,
            credentials=creds,
            db_session=db,
            orchestrator=mock_orch,
            drive_client=mock_drive,
        )

    db.refresh(job)
    assert job.status == JobStatus.COMPLETED
    assert job.result is not None
    assert job.result["session"] == 105
    assert job.result["slo"] == 1
    assert "1051" in job.result["original_file"]

    dl_kwargs = mock_orch.download_worksheet.call_args[1]
    assert dl_kwargs["filename"] == "1051.docx"

    sub_kwargs = mock_orch.submit_worksheet_link.call_args[1]
    assert sub_kwargs["session"] == 105
    assert sub_kwargs["slo"] == 1


@pytest.mark.parametrize("ws_id_input", [
    "1052",
    "105_2",
    "105-2",
    "1052.docx",
    "Unit 1 Session 5 SLO 2",
    "Unit 1 - Session 5 - SLO 2",
    "Session 5 SLO 2",
])
@pytest.mark.asyncio
async def test_worksheet_id_parsing_formats_for_slo2(db_session: Session, ws_id_input: str):
    """Verify various string formats for worksheet_id all resolve correctly to session=105, slo=2."""
    db = db_session
    job = Job(
        user_id="RA2111003010001",
        course_id="21CSC303J",
        semester_id="3",
        worksheet_id=ws_id_input,
        transport_mode="http",
        status=JobStatus.PENDING,
    )
    db.add(job)
    db.commit()
    db.refresh(job)

    mock_orch = _setup_mock_orchestrator_for_session(105)
    mock_drive = _setup_mock_drive()

    creds = {"USER_ID": "RA2111003010001", "PASSWORD": "secretpassword"}

    with patch("apps.worker.tasks.store_job_credentials"), \
         patch("apps.worker.tasks.get_job_credentials", return_value=creds), \
         patch("apps.worker.tasks.clear_job_credentials"):
        await _run_job_workflow(
            job_id=job.id,
            credentials=creds,
            db_session=db,
            orchestrator=mock_orch,
            drive_client=mock_drive,
        )

        db.refresh(job)
        assert job.status == JobStatus.AWAITING_USER_REVIEW

        await _run_submission_workflow(
            job_id=job.id,
            credentials=creds,
            db_session=db,
            orchestrator=mock_orch,
            drive_client=mock_drive,
        )

    db.refresh(job)
    assert job.status == JobStatus.COMPLETED
    assert job.result["slo"] == 2
    dl_kwargs = mock_orch.download_worksheet.call_args[1]
    assert dl_kwargs["filename"] == "1052.docx"


@pytest.mark.asyncio
async def test_slo2_idempotency_does_not_falsely_skip_when_slo1_is_verified(db_session: Session):
    """Verify that if SLO 1 is already verified on SRM (1051: 2), SLO 2 (1052: 0) still gets submitted."""
    db = db_session
    job = Job(
        user_id="RA2111003010001",
        course_id="21CSC303J",
        semester_id="3",
        worksheet_id="1052",
        transport_mode="http",
        status=JobStatus.PENDING,
    )
    db.add(job)
    db.commit()
    db.refresh(job)

    mock_orch = _setup_mock_orchestrator_for_session(105)
    # SLO 1 is verified (2), but SLO 2 is not submitted (0)
    mock_orch.get_session_status = AsyncMock(
        return_value=SRMSessionStatus(
            session=105,
            practice_status={"1051": 2, "1052": 0},
            slo_links={"1051": "https://docs.google.com/document/d/1A2B3C4D5E6F7G8H9I0J_slo1/edit"},
        )
    )
    mock_drive = _setup_mock_drive()

    creds = {
        "USER_ID": "RA2111003010001",
        "PASSWORD": "secretpassword",
        "requested_session": 105,
        "requested_slo": 2,
    }

    with patch("apps.worker.tasks.store_job_credentials"), \
         patch("apps.worker.tasks.get_job_credentials", return_value=creds), \
         patch("apps.worker.tasks.clear_job_credentials"):
        await _run_job_workflow(
            job_id=job.id,
            credentials=creds,
            db_session=db,
            orchestrator=mock_orch,
            drive_client=mock_drive,
        )

        db.refresh(job)
        assert job.status == JobStatus.AWAITING_USER_REVIEW

        await _run_submission_workflow(
            job_id=job.id,
            credentials=creds,
            db_session=db,
            orchestrator=mock_orch,
            drive_client=mock_drive,
        )

    db.refresh(job)
    assert job.status == JobStatus.COMPLETED

    # Crucial check: submit_worksheet_link was called for SLO 2!
    mock_orch.submit_worksheet_link.assert_called_once()
    sub_kwargs = mock_orch.submit_worksheet_link.call_args[1]
    assert sub_kwargs["session"] == 105
    assert sub_kwargs["slo"] == 2


@pytest.mark.asyncio
async def test_verify_submission_candidate_keys_target_specific_slo():
    """Verify SRMHttpClient.verify_submission targets specific SLO key, not other SLOs."""
    client = SRMHttpClient(base_url="https://srm.portal")

    # Mock get_session_status returning distinct links for SLO 1 and SLO 2
    mock_status = SRMSessionStatus(
        session=105,
        practice_status={"1051": 2, "1052": 2},
        slo_links={
            "1051": "https://docs.google.com/document/d/1A2B3C4D5E6F7G8H9I0J_slo1/edit",
            "1052": "https://docs.google.com/document/d/1A2B3C4D5E6F7G8H9I0J_slo2/edit",
        },
    )
    client.get_session_status = AsyncMock(return_value=mock_status)

    course_info = {"BATCH_ID": "B1", "COURSE_CODE": "21CSC303J"}

    # Verifying SLO 2 with SLO 2 link should succeed
    ok_slo2 = await client.verify_submission(
        session_or_worksheet_id=105,
        slo=2,
        expected_link="https://docs.google.com/document/d/1A2B3C4D5E6F7G8H9I0J_slo2/view",
        course_info=course_info,
    )
    assert ok_slo2 is True

    # Verifying SLO 2 with SLO 1 link should fail (does not mistakenly match SLO 1)
    from packages.srm.exceptions import VerificationFailed
    with pytest.raises(VerificationFailed):
        await client.verify_submission(
            session_or_worksheet_id=105,
            slo=2,
            expected_link="https://docs.google.com/document/d/1A2B3C4D5E6F7G8H9I0J_slo1/view",
            course_info=course_info,
        )
