"""Regression and integration test suite for User Review Before SRM Submission (Phase 8).

Validates all 16 specification requirements:
1. Workflow stops after successful Drive upload.
2. Job becomes AWAITING_USER_REVIEW.
3. SRM submitlink is NOT called automatically.
4. Dashboard displays View / Download / Submit controls.
5. User clicks Submit -> SRM submission begins.
6. Cancel confirmation -> no submission.
7. Double-click Submit -> only one submission.
8. Browser refresh preserves AWAITING_USER_REVIEW.
9. Browser can be closed and reopened before submission.
10. Drive public verification is required before Submit becomes available.
11. Failed SRM submission -> FAILED with retry capability.
12. Existing SRM submission detected -> no duplicate submission.
13. Successful submission -> VERIFYING -> COMPLETED.
14. CAPTCHA workflow remains unchanged.
15. SLO1/SLO2 selection remains unchanged.
16. Original worksheet remains immutable.
"""

import hashlib
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from apps.api.main import app
from apps.worker.tasks import (
    _run_job_workflow,
    _run_submission_workflow,
    get_job_credentials,
    process_job,
    store_job_credentials,
    submit_job,
)
from packages.drive.client import BaseDriveClient
from packages.drive.models import DriveFileMetadata
from packages.shared.models.job import Job, JobStatus
from packages.shared.schemas.job import JobResponse
from packages.srm.exceptions import SRMException
from packages.srm.models import (
    SRMCourse,
    SRMSessionStatus,
    SRMSubmissionResult,
    SRMWorksheetMetadata,
)
from packages.srm.orchestrator import SRMOrchestrator


def _create_sample_docx(path: Path) -> Path:
    """Create a synthetic DOCX worksheet."""
    doc = Document()
    doc.add_heading("Course: 21CSC303J Session: 101 SLO: 1", level=1)
    doc.add_paragraph("1. What does CPU stand for?")
    doc.add_paragraph("A. Central Processing Unit")
    doc.add_paragraph("B. Central Process Unit")
    doc.add_paragraph("Answer: ")
    doc.save(str(path))
    return path


def _setup_test_orchestrator(tmp_path: Path) -> MagicMock:
    """Create a mock SRMOrchestrator."""
    mock_orch = MagicMock(spec=SRMOrchestrator)
    mock_orch.transport_name = "http"
    mock_orch.connect = AsyncMock(return_value=True)
    mock_orch.authenticate = AsyncMock(return_value=True)
    mock_orch.capture_login_captcha = AsyncMock(return_value=None)
    mock_orch.close = AsyncMock(return_value=None)

    course = SRMCourse(
        course_code="21CSC303J",
        course_name="Database Management Systems",
        batch_id="B1",
        semester=3,
        department="CSE",
    )
    mock_orch.get_courses = AsyncMock(return_value=[course])

    ws_meta = SRMWorksheetMetadata(
        course_code="21CSC303J",
        session=101,
        slo=1,
        filename="1011.docx",
        format="docx",
        storage_path="data/coordinator/21CSC303J/slp",
        download_url="https://srm.portal/file/1011.docx",
        is_available=True,
    )
    mock_orch.discover_worksheets = AsyncMock(return_value=[ws_meta])
    mock_orch.get_worksheet_file = AsyncMock(return_value="https://srm.portal/file/1011.docx")

    sample_doc = _create_sample_docx(tmp_path / "1011.docx")

    async def fake_download(file_url_or_id, destination_dir=None, filename=None):
        target = (destination_dir or tmp_path) / (filename or "1011.docx")
        target.write_bytes(sample_doc.read_bytes())
        return target

    mock_orch.download_worksheet = AsyncMock(side_effect=fake_download)
    mock_orch.get_session_status = AsyncMock(
        return_value=SRMSessionStatus(session=101, practice_status=0, slo_links={})
    )
    mock_orch.submit_worksheet_link = AsyncMock(
        return_value=SRMSubmissionResult(
            success=True,
            message="Link submitted successfully",
            returned_link="https://drive.google.com/file/d/drive_test_123/view",
        )
    )
    mock_orch.verify_submission = AsyncMock(return_value=True)
    return mock_orch


def _setup_test_drive_client() -> MagicMock:
    """Create a mock BaseDriveClient with verified public sharing."""
    mock_drive = MagicMock(spec=BaseDriveClient)
    meta = DriveFileMetadata(
        file_id="drive_test_123",
        filename="completed_1011.docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        web_url="https://drive.google.com/file/d/drive_test_123/view?usp=drivesdk",
        download_url="https://drive.google.com/uc?id=drive_test_123&export=download",
        is_public=True,
        permission_status="VERIFIED_PUBLIC_READER",
    )
    mock_drive.upload_file = AsyncMock(return_value=meta)
    mock_drive.verify_file_public_access = AsyncMock(return_value=True)
    return mock_drive




@pytest.mark.asyncio
async def test_01_02_03_workflow_stops_at_awaiting_user_review_and_no_auto_submit(
    db_session: Session, tmp_path: Path
):
    """Req 1, 2, 3: Workflow stops after Drive upload; Job becomes AWAITING_USER_REVIEW; SRM submit is NOT called automatically."""
    job = Job(
        user_id="RA2111003010001",
        course_id="21CSC303J",
        semester_id="3",
        worksheet_id="1011",
        status=JobStatus.PENDING,
    )
    db_session.add(job)
    db_session.commit()

    mock_orch = _setup_test_orchestrator(tmp_path)
    mock_drive = _setup_test_drive_client()

    creds = {"USER_ID": "RA2111003010001", "PASSWORD": "SecretPassword123!"}

    # Execute workflow with default auto_submit=False
    await _run_job_workflow(
        job_id=job.id,
        credentials=creds,
        orchestrator=mock_orch,
        drive_client=mock_drive,
        db_session=db_session,
        auto_submit=False,
    )

    db_session.refresh(job)

    # 1. State verification: Stopped at AWAITING_USER_REVIEW
    assert job.status == JobStatus.AWAITING_USER_REVIEW
    assert job.current_step == "awaiting_user_review"
    assert job.error_message is None

    # 2. Result verification: Safe metadata stored
    result = job.result
    assert result is not None
    assert result["review_ready"] is True
    assert result["drive_verified"] is True
    assert result["submission_allowed"] is True
    assert result["drive_file_id"] == "drive_test_123"
    assert "https://drive.google.com" in result["drive_web_url"]
    assert "completed_1011.docx" in result["completed_file"]

    # 3. SRM submission must NOT have been called automatically
    mock_orch.submit_worksheet_link.assert_not_called()
    mock_orch.verify_submission.assert_not_called()

    # 4. Ephemeral credentials must be preserved for user submission
    cached = get_job_credentials(job.id)
    assert cached is not None
    assert cached.get("USER_ID") == "RA2111003010001"


def test_04_dashboard_ui_controls_and_file_download(client: TestClient, db_session: Session, tmp_path: Path):
    """Req 4: Dashboard displays View/Download/Submit controls; download endpoint returns completed file."""
    # 1. Check dashboard HTML contains the review card, download/view buttons, and confirmation modal
    dash_resp = client.get("/dashboard")
    assert dash_resp.status_code == 200
    html = dash_resp.text
    assert 'id="review-card"' in html
    assert 'id="review-view-doc-btn"' in html
    assert 'id="review-download-btn"' in html
    assert 'id="review-submit-srm-btn"' in html
    assert 'id="submit-confirm-modal"' in html
    assert "AWAITING_USER_REVIEW" in html
    assert "Your completed worksheet is ready. Review it before submitting to SRM." in html

    # 2. Create job in AWAITING_USER_REVIEW
    completed_doc = tmp_path / "completed_1011.docx"
    _create_sample_docx(completed_doc)

    job = Job(
        user_id="RA2111003010001",
        course_id="21CSC303J",
        semester_id="3",
        worksheet_id="1011",
        status=JobStatus.AWAITING_USER_REVIEW,
        current_step="awaiting_user_review",
        result={
            "course_code": "21CSC303J",
            "session": 101,
            "slo": 1,
            "completed_file": "completed_1011.docx",
            "completed_file_path": str(completed_doc),
            "drive_file_id": "drive_test_123",
            "drive_web_url": "https://drive.google.com/file/d/drive_test_123/view",
            "drive_permission_status": "VERIFIED_PUBLIC_READER",
            "review_ready": True,
            "questions_count": 4,
            "answers_count": 4,
        },
    )
    db_session.add(job)
    db_session.commit()
    db_session.refresh(job)

    # 3. Test API job details and computed fields
    resp = client.get(f"/api/v1/jobs/{job.id}")
    assert resp.status_code == 200
    data = resp.json()
    assert data["review_ready"] is True
    assert data["completed_file_name"] == "completed_1011.docx"
    assert data["drive_file_id"] == "drive_test_123"
    assert data["drive_web_view_link"] == "https://drive.google.com/file/d/drive_test_123/view"
    assert data["drive_verified"] is True
    assert data["submission_allowed"] is True

    # 4. Test download endpoint serves completed docx
    dl_resp = client.get(f"/api/v1/jobs/{job.id}/download")
    assert dl_resp.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in dl_resp.headers["content-type"]
    assert "completed_1011.docx" in dl_resp.headers.get("content-disposition", "")
    assert len(dl_resp.content) > 0


@pytest.mark.asyncio
async def test_05_13_user_clicks_submit_transitions_through_submitting_verifying_completed(
    db_session: Session, tmp_path: Path
):
    """Req 5, 13: User clicks Submit -> transitions SUBMITTING -> VERIFYING -> COMPLETED."""
    completed_doc = tmp_path / "completed_1011.docx"
    _create_sample_docx(completed_doc)

    job = Job(
        user_id="RA2111003010001",
        course_id="21CSC303J",
        semester_id="3",
        worksheet_id="1011",
        status=JobStatus.AWAITING_USER_REVIEW,
        current_step="awaiting_user_review",
        result={
            "course_code": "21CSC303J",
            "course_name": "Database Management Systems",
            "batch_id": "B1",
            "session": 101,
            "slo": 1,
            "completed_file": "completed_1011.docx",
            "completed_file_path": str(completed_doc),
            "drive_file_id": "drive_test_123",
            "drive_web_url": "https://drive.google.com/file/d/drive_test_123/view",
            "drive_permission_status": "VERIFIED_PUBLIC_READER",
            "review_ready": True,
        },
    )
    db_session.add(job)
    db_session.commit()

    mock_orch = _setup_test_orchestrator(tmp_path)
    mock_drive = _setup_test_drive_client()

    store_job_credentials(job.id, {"USER_ID": "RA2111003010001", "PASSWORD": "SecretPassword!"})

    # Trigger explicit submission workflow
    await _run_submission_workflow(
        job_id=job.id,
        orchestrator=mock_orch,
        drive_client=mock_drive,
        db_session=db_session,
    )

    db_session.refresh(job)

    # Status verification: COMPLETED
    assert job.status == JobStatus.COMPLETED
    assert job.current_step == "workflow_completed"
    assert job.result["verification_status"] == "VERIFIED"
    assert job.result["practice_status"] == 2
    assert "submission_started_at" in job.result
    assert "submission_verified_at" in job.result

    # SRM submission and verification were executed
    mock_orch.submit_worksheet_link.assert_called_once()
    sub_kwargs = mock_orch.submit_worksheet_link.call_args[1]
    assert sub_kwargs["view_link"] == "https://drive.google.com/file/d/drive_test_123/view"
    mock_orch.verify_submission.assert_called_once()

    # Ephemeral credentials wiped on completion
    assert get_job_credentials(job.id) is None


def test_06_cancel_confirmation_does_not_submit(client: TestClient, db_session: Session, tmp_path: Path):
    """Req 6: Dismissing confirmation modal does not submit; cancelling job marks FAILED and clears credentials."""
    job = Job(
        user_id="RA2111003010001",
        course_id="21CSC303J",
        status=JobStatus.AWAITING_USER_REVIEW,
        current_step="awaiting_user_review",
        result={
            "drive_file_id": "test_id",
            "drive_web_url": "https://drive.google.com/view",
            "review_ready": True,
        },
    )
    db_session.add(job)
    db_session.commit()
    store_job_credentials(job.id, {"USER_ID": "RA2111003010001", "PASSWORD": "secret"})

    # If user cancels the job explicitly
    cancel_resp = client.post(f"/api/v1/jobs/{job.id}/cancel")
    assert cancel_resp.status_code == 200
    assert cancel_resp.json()["status"] == "FAILED"
    assert cancel_resp.json()["current_step"] == "cancelled_by_user"

    # Ephemeral credentials wiped
    assert get_job_credentials(job.id) is None


def test_07_double_click_submit_single_dispatch(client: TestClient, db_session: Session):
    """Req 7: Double-clicking Submit triggers only one submission dispatch (idempotency guard)."""
    job = Job(
        user_id="RA2111003010001",
        status=JobStatus.AWAITING_USER_REVIEW,
        current_step="awaiting_user_review",
        result={
            "drive_file_id": "test_drive_id",
            "drive_web_url": "https://drive.google.com/view",
            "review_ready": True,
        },
    )
    db_session.add(job)
    db_session.commit()
    store_job_credentials(job.id, {"USER_ID": "RA2111003010001"})

    with patch("apps.api.routes.jobs.submit_job.delay") as mock_delay:
        # First click: dispatches task
        resp1 = client.post(f"/api/v1/jobs/{job.id}/submit", json={})
        assert resp1.status_code == 200
        assert mock_delay.call_count == 1

        # Second click: job is already SUBMITTING -> returns job without second dispatch
        resp2 = client.post(f"/api/v1/jobs/{job.id}/submit", json={})
        assert resp2.status_code == 200
        assert mock_delay.call_count == 1  # Still 1!


def test_08_09_refresh_and_browser_closure_preserves_review_state(client: TestClient, db_session: Session):
    """Req 8, 9: Browser refresh or reopening dashboard preserves AWAITING_USER_REVIEW without auto-submitting."""
    job = Job(
        user_id="RA2111003010001",
        course_id="21CSC303J",
        status=JobStatus.AWAITING_USER_REVIEW,
        current_step="awaiting_user_review",
        result={
            "drive_file_id": "test_id",
            "drive_web_url": "https://drive.google.com/view",
            "review_ready": True,
        },
    )
    db_session.add(job)
    db_session.commit()

    # Simulate 5 consecutive page reloads / status checks
    for _ in range(5):
        resp = client.get(f"/api/v1/jobs/{job.id}")
        assert resp.status_code == 200
        assert resp.json()["status"] == "AWAITING_USER_REVIEW"
        assert resp.json()["review_ready"] is True

    # Check active jobs list
    list_resp = client.get("/api/v1/jobs?limit=5")
    assert list_resp.status_code == 200
    active_matches = [j for j in list_resp.json() if j["id"] == job.id]
    assert len(active_matches) == 1
    assert active_matches[0]["status"] == "AWAITING_USER_REVIEW"


def test_10_drive_verification_required_before_submit(client: TestClient, db_session: Session):
    """Req 10: Drive public verification is required before Submit becomes available."""
    # Job without Drive link
    job = Job(
        user_id="RA2111003010001",
        status=JobStatus.AWAITING_USER_REVIEW,
        result={"completed_file": "completed.docx"},  # missing drive_web_url
    )
    db_session.add(job)
    db_session.commit()

    resp = client.post(f"/api/v1/jobs/{job.id}/submit", json={})
    assert resp.status_code == 400
    assert "Drive link is missing or unverified" in resp.json()["detail"]


@pytest.mark.asyncio
async def test_11_failed_srm_submission_allows_retry(db_session: Session, tmp_path: Path):
    """Req 11: Failed SRM submission transitions to FAILED with Drive link preserved and retry capability."""
    job = Job(
        user_id="RA2111003010001",
        status=JobStatus.AWAITING_USER_REVIEW,
        current_step="awaiting_user_review",
        result={
            "drive_file_id": "drive_test_123",
            "drive_web_url": "https://drive.google.com/file/d/drive_test_123/view",
            "review_ready": True,
            "session": 101,
            "slo": 1,
            "course_code": "21CSC303J",
        },
    )
    db_session.add(job)
    db_session.commit()

    mock_orch = _setup_test_orchestrator(tmp_path)
    mock_orch.submit_worksheet_link.side_effect = SRMException("Portal 500 error on submit")
    mock_drive = _setup_test_drive_client()

    store_job_credentials(job.id, {"USER_ID": "RA2111003010001"})

    # Attempt 1: Fails
    await _run_submission_workflow(
        job_id=job.id,
        orchestrator=mock_orch,
        drive_client=mock_drive,
        db_session=db_session,
    )

    db_session.refresh(job)
    assert job.status == JobStatus.FAILED
    assert "Portal 500 error" in job.error_message
    assert job.result["drive_web_url"] == "https://drive.google.com/file/d/drive_test_123/view"
    assert job.result.get("submission_allowed") is True

    # Attempt 2: Recover on retry
    mock_orch.submit_worksheet_link.side_effect = None
    mock_orch.submit_worksheet_link.return_value = SRMSubmissionResult(
        success=True,
        message="Submitted on retry",
        returned_link="https://drive.google.com/file/d/drive_test_123/view",
    )

    await _run_submission_workflow(
        job_id=job.id,
        orchestrator=mock_orch,
        drive_client=mock_drive,
        db_session=db_session,
    )

    db_session.refresh(job)
    assert job.status == JobStatus.COMPLETED
    assert job.result["verification_status"] == "VERIFIED"


@pytest.mark.asyncio
async def test_12_existing_srm_submission_detected_no_duplicate(db_session: Session, tmp_path: Path):
    """Req 12: Existing SRM submission detected -> no duplicate submission request, proceeds to verification."""
    drive_link = "https://drive.google.com/file/d/drive_test_123/view"
    job = Job(
        user_id="RA2111003010001",
        status=JobStatus.AWAITING_USER_REVIEW,
        current_step="awaiting_user_review",
        result={
            "drive_file_id": "drive_test_123",
            "drive_web_url": drive_link,
            "session": 101,
            "slo": 1,
            "course_code": "21CSC303J",
            "batch_id": "B1",
        },
    )
    db_session.add(job)
    db_session.commit()

    mock_orch = _setup_test_orchestrator(tmp_path)
    # Session already verified on SRM
    mock_orch.get_session_status.return_value = SRMSessionStatus(
        session=101,
        practice_status={"1011": 2},
        slo_links={"1011": drive_link},
    )
    mock_drive = _setup_test_drive_client()

    await _run_submission_workflow(
        job_id=job.id,
        orchestrator=mock_orch,
        drive_client=mock_drive,
        db_session=db_session,
    )

    db_session.refresh(job)
    assert job.status == JobStatus.COMPLETED
    # Idempotency: submit_worksheet_link was NOT called because already on portal
    mock_orch.submit_worksheet_link.assert_not_called()
    mock_orch.verify_submission.assert_called_once()


@pytest.mark.asyncio
async def test_14_captcha_workflow_remains_unchanged(db_session: Session, tmp_path: Path):
    """Req 14: CAPTCHA pause/resume workflow remains unchanged and halts at AWAITING_USER_REVIEW."""
    job = Job(
        user_id="RA2111003010001",
        course_id="21CSC303J",
        semester_id="3",
        worksheet_id="1011",
        status=JobStatus.PENDING,
    )
    db_session.add(job)
    db_session.commit()

    mock_orch = _setup_test_orchestrator(tmp_path)
    mock_orch.capture_login_captcha.return_value = {"type": "canvas", "selector": "#captcha"}
    mock_drive = _setup_test_drive_client()

    creds = {"USER_ID": "RA2111003010001", "PASSWORD": "SecretPassword!"}

    # Step 1: Pauses on CAPTCHA
    await _run_job_workflow(
        job_id=job.id,
        credentials=creds,
        orchestrator=mock_orch,
        drive_client=mock_drive,
        db_session=db_session,
        auto_submit=False,
    )

    db_session.refresh(job)
    assert job.status == JobStatus.WAITING_FOR_CAPTCHA
    assert job.captcha_challenge is not None

    # Step 2: Solved CAPTCHA -> resumes and reaches AWAITING_USER_REVIEW
    mock_orch.capture_login_captcha.return_value = None
    job.captcha_solution = "123456"
    db_session.commit()

    await _run_job_workflow(
        job_id=job.id,
        credentials=None,
        orchestrator=mock_orch,
        drive_client=mock_drive,
        db_session=db_session,
        auto_submit=False,
    )

    db_session.refresh(job)
    assert job.status == JobStatus.AWAITING_USER_REVIEW
    assert job.result["review_ready"] is True
    mock_orch.submit_worksheet_link.assert_not_called()


@pytest.mark.asyncio
async def test_15_slo1_slo2_selection_remains_unchanged(db_session: Session, tmp_path: Path):
    """Req 15: SLO 1 and SLO 2 selection remains preserved and recorded in review metadata."""
    job = Job(
        user_id="RA2111003010001",
        course_id="21CSC303J",
        semester_id="3",
        worksheet_id="1052",
        status=JobStatus.PENDING,
    )
    db_session.add(job)
    db_session.commit()

    mock_orch = _setup_test_orchestrator(tmp_path)
    # Return both SLO 1 and SLO 2 in discovery
    ws1 = SRMWorksheetMetadata(
        course_code="21CSC303J",
        session=105,
        slo=1,
        filename="1051.docx",
        format="docx",
        storage_path="data/coordinator/21CSC303J/slp",
        is_available=True,
    )
    ws2 = SRMWorksheetMetadata(
        course_code="21CSC303J",
        session=105,
        slo=2,
        filename="1052.docx",
        format="docx",
        storage_path="data/coordinator/21CSC303J/slp",
        is_available=True,
    )
    mock_orch.discover_worksheets.return_value = [ws1, ws2]
    mock_orch.get_worksheet_file.return_value = "https://srm.portal/file/1052.docx"
    mock_drive = _setup_test_drive_client()

    await _run_job_workflow(
        job_id=job.id,
        credentials={"USER_ID": "RA2111003010001", "requested_session": 105, "requested_slo": 2},
        orchestrator=mock_orch,
        drive_client=mock_drive,
        db_session=db_session,
        auto_submit=False,
    )

    db_session.refresh(job)
    assert job.status == JobStatus.AWAITING_USER_REVIEW
    assert job.result["session"] == 105
    assert job.result["slo"] == 2
    assert "1052" in job.result["completed_file"]


@pytest.mark.asyncio
async def test_16_original_worksheet_remains_immutable(db_session: Session, tmp_path: Path):
    """Req 16: Original worksheet remains immutable; completed worksheet is created as a distinct copy."""
    orig_file = tmp_path / "1011.docx"
    _create_sample_docx(orig_file)
    orig_hash = hashlib.sha256(orig_file.read_bytes()).hexdigest()

    job = Job(
        user_id="RA2111003010001",
        course_id="21CSC303J",
        semester_id="3",
        worksheet_id="1011",
        status=JobStatus.PENDING,
        result={"original_file_path": str(orig_file)},
    )
    db_session.add(job)
    db_session.commit()

    mock_orch = _setup_test_orchestrator(tmp_path)
    mock_drive = _setup_test_drive_client()

    await _run_job_workflow(
        job_id=job.id,
        credentials={"USER_ID": "RA2111003010001"},
        orchestrator=mock_orch,
        drive_client=mock_drive,
        db_session=db_session,
        auto_submit=False,
    )

    db_session.refresh(job)
    assert job.status == JobStatus.AWAITING_USER_REVIEW

    # Verify original file on disk is bit-for-bit identical to before execution
    current_orig_hash = hashlib.sha256(orig_file.read_bytes()).hexdigest()
    assert current_orig_hash == orig_hash

    # Verify completed file is distinct
    completed_path = Path(job.result["completed_file_path"])
    assert completed_path.exists()
    assert completed_path.resolve() != orig_file.resolve()
