"""Tests for Answer Generation, Document Filling, and End-to-End Pipeline (Milestone 5).

Verifies:
1. Complete pipeline execution on MCQ worksheets.
2. Complete pipeline execution on One-Word / Fill-in-the-blank worksheets.
3. Complete pipeline execution on Short Answer worksheets.
4. Complete pipeline execution on Long Answer worksheets.
5. Complete pipeline execution on Mixed worksheets.
6. Table-based worksheet filling.
7. Original document immutability (non-destructive guarantee).
8. Answer engine replaceability (custom provider injection).
9. Error and low-confidence handling.
10. End-to-end processing of real SRM 1011.docx worksheet.
11. DefaultWorksheetProcessor adapter integration.
"""

import hashlib
from pathlib import Path
import docx
import pytest

from packages.worksheets.answer_engine import (
    AnswerEngineFactory,
    BaseAnswerEngine,
    RuleBasedAnswerEngine,
)
from packages.worksheets.answer_models import (
    AnswerStatus,
    GeneratedAnswer,
    WorksheetAnswers,
)
from packages.worksheets.filler import DocxWorksheetFiller
from packages.worksheets.models import ParsedQuestion, ParsedWorksheet, QuestionType
from packages.worksheets.pipeline import WorksheetPipeline
from packages.worksheets.processor import DefaultWorksheetProcessor

# Reuse test helpers from test_worksheet_parser
from tests.test_worksheet_parser import (
    _create_synthetic_long_answer_docx,
    _create_synthetic_mcq_docx,
    _create_synthetic_mixed_docx,
    _create_synthetic_one_word_docx,
    _create_synthetic_short_answer_docx,
    _create_synthetic_table_docx,
)


@pytest.mark.asyncio
async def test_mcq_answer_generation_and_filling(tmp_path: Path):
    """Verify end-to-end MCQ parsing, option selection, and document filling."""
    orig_file = _create_synthetic_mcq_docx(tmp_path / "mcq_input.docx")
    pipeline = WorksheetPipeline()

    result = await pipeline.process(orig_file, output_dir=tmp_path)

    assert result.success is True
    assert result.completed_file.exists()
    assert result.completed_file != orig_file
    assert result.completed_file.name == "completed_mcq_input.docx"

    # Verify answers generated
    answers = result.answers
    assert answers.total_count == 4
    for ans in answers.answers:
        assert ans.status == AnswerStatus.SUCCESS
        assert ans.selected_option in ("A", "B", "C", "D")
        assert ans.confidence >= 0.7

    # Verify completed document contents
    comp_doc = docx.Document(str(result.completed_file))
    full_text = "\n".join(p.text for p in comp_doc.paragraphs)
    assert "Answer: " in full_text
    assert "Central Processing Unit" in full_text


@pytest.mark.asyncio
async def test_one_word_answer_generation_and_filling(tmp_path: Path):
    """Verify one-word and fill-in-the-blank answer generation and filling."""
    orig_file = _create_synthetic_one_word_docx(tmp_path / "one_word_input.docx")
    pipeline = WorksheetPipeline()

    result = await pipeline.process(orig_file, output_dir=tmp_path)

    assert result.success is True
    assert result.answers.total_count == 4
    for ans in result.answers.answers:
        assert ans.status == AnswerStatus.SUCCESS
        assert len(ans.answer_text) > 0
        assert ans.confidence >= 0.8

    # Verify completed document contains answers
    comp_doc = docx.Document(str(result.completed_file))
    full_text = "\n".join(p.text for p in comp_doc.paragraphs)
    assert "Answer: " in full_text
    assert "Compilation" in full_text or "Compiler" in full_text


@pytest.mark.asyncio
async def test_short_answer_generation_and_filling(tmp_path: Path):
    """Verify concise conceptual answer generation for 1-4 marks questions."""
    orig_file = _create_synthetic_short_answer_docx(tmp_path / "short_input.docx")
    pipeline = WorksheetPipeline()

    result = await pipeline.process(orig_file, output_dir=tmp_path)

    assert result.success is True
    assert result.answers.total_count == 4

    comp_doc = docx.Document(str(result.completed_file))
    full_text = "\n".join(p.text for p in comp_doc.paragraphs)
    assert "Answer: " in full_text
    assert "polymorphism" in full_text.lower()


@pytest.mark.asyncio
async def test_long_answer_generation_and_filling(tmp_path: Path):
    """Verify detailed multi-section answers for descriptive questions."""
    orig_file = _create_synthetic_long_answer_docx(tmp_path / "long_input.docx")
    pipeline = WorksheetPipeline()

    result = await pipeline.process(orig_file, output_dir=tmp_path)

    assert result.success is True
    assert result.answers.total_count == 3

    comp_doc = docx.Document(str(result.completed_file))
    full_text = "\n".join(p.text for p in comp_doc.paragraphs)
    assert "Architectural" in full_text or "Architecture" in full_text
    assert "1." in full_text


@pytest.mark.asyncio
async def test_mixed_worksheet_end_to_end(tmp_path: Path):
    """Verify processing across mixed sections (MCQ, One-Word, Short, Long)."""
    orig_file = _create_synthetic_mixed_docx(tmp_path / "mixed_input.docx")
    pipeline = WorksheetPipeline()

    result = await pipeline.process(orig_file, output_dir=tmp_path)

    assert result.success is True
    assert result.answers.total_count == 4

    # Verify each question type was answered
    mcq_ans = result.answers.answers[0]
    assert mcq_ans.question_type == QuestionType.MCQ
    assert mcq_ans.selected_option is not None

    one_word_ans = result.answers.answers[1]
    assert one_word_ans.question_type == QuestionType.ONE_WORD

    short_ans = result.answers.answers[2]
    assert short_ans.question_type == QuestionType.SHORT_ANSWER

    long_ans = result.answers.answers[3]
    assert long_ans.question_type == QuestionType.LONG_ANSWER


@pytest.mark.asyncio
async def test_table_worksheet_filling(tmp_path: Path):
    """Verify filling answers inside table rows."""
    orig_file = _create_synthetic_table_docx(tmp_path / "table_input.docx")
    pipeline = WorksheetPipeline()

    result = await pipeline.process(orig_file, output_dir=tmp_path)

    assert result.success is True
    assert result.answers.total_count == 3

    comp_doc = docx.Document(str(result.completed_file))
    table = comp_doc.tables[0]
    # Check that answers were injected into table rows
    all_table_text = " ".join(c.text for row in table.rows for c in row.cells)
    assert "Answer:" in all_table_text


@pytest.mark.asyncio
async def test_original_document_immutability(tmp_path: Path):
    """Verify that original file bytes remain 100% unmodified during filling."""
    orig_file = _create_synthetic_mixed_docx(tmp_path / "immutable_test.docx")
    hash_before = hashlib.sha256(orig_file.read_bytes()).hexdigest()

    pipeline = WorksheetPipeline()
    result = await pipeline.process(orig_file, output_dir=tmp_path)

    hash_after = hashlib.sha256(orig_file.read_bytes()).hexdigest()
    assert hash_before == hash_after, "Original document was modified!"
    assert result.completed_file.stat().st_size > 0


@pytest.mark.asyncio
async def test_answer_engine_replaceability(tmp_path: Path):
    """Verify the answer engine can be replaced with a custom provider seamlessly."""
    orig_file = _create_synthetic_short_answer_docx(tmp_path / "custom_provider_test.docx")

    class CustomMockProvider(BaseAnswerEngine):
        async def generate_answers(self, worksheet, context=None):
            ans_list = [
                GeneratedAnswer(
                    question_id=q.question_id,
                    question_number=q.question_number,
                    question_type=q.question_type,
                    answer_text=f"CUSTOM_AI_SOLUTION_FOR_{q.question_id}",
                    confidence=0.99,
                )
                for q in worksheet.questions
            ]
            return WorksheetAnswers(
                worksheet_filename=worksheet.filename,
                answers=ans_list,
                provider="custom_mock",
            )

        async def answer_question(self, question, worksheet_context=None):
            return GeneratedAnswer(
                question_id=question.question_id,
                answer_text="CUSTOM_AI_SOLUTION",
            )

    custom_pipeline = WorksheetPipeline(answer_engine=CustomMockProvider())
    result = await custom_pipeline.process(orig_file, output_dir=tmp_path)

    assert result.success is True
    assert result.answers.provider == "custom_mock"
    comp_doc = docx.Document(str(result.completed_file))
    full_text = "\n".join(p.text for p in comp_doc.paragraphs)
    assert "CUSTOM_AI_SOLUTION_FOR_" in full_text


@pytest.mark.asyncio
async def test_error_and_low_confidence_handling(tmp_path: Path):
    """Verify low-confidence and unparseable question handling."""
    engine = RuleBasedAnswerEngine()

    # Question with empty prompt
    bad_question = ParsedQuestion(
        question_id="bad_1",
        question_text="",
        question_type=QuestionType.UNKNOWN,
    )
    ans = await engine.answer_question(bad_question)
    assert ans.status == AnswerStatus.LOW_CONFIDENCE
    assert ans.confidence <= 0.3

    # MCQ with missing options
    mcq_no_opts = ParsedQuestion(
        question_id="bad_mcq",
        question_text="Which of the following is correct?",
        question_type=QuestionType.MCQ,
        options=[],
    )
    mcq_ans = await engine.answer_question(mcq_no_opts)
    assert mcq_ans.status == AnswerStatus.LOW_CONFIDENCE
    assert mcq_ans.confidence <= 0.3


@pytest.mark.asyncio
async def test_real_1011_docx_end_to_end_pipeline(tmp_path: Path):
    """Verify end-to-end answer generation and filling on real SRM 1011.docx."""
    real_path = Path("artifacts/real_1011.docx")
    if not real_path.exists():
        pytest.skip("Real 1011.docx not found, skipping.")

    orig_hash = hashlib.sha256(real_path.read_bytes()).hexdigest()

    pipeline = WorksheetPipeline()
    result = await pipeline.process(
        worksheet_path=real_path,
        output_dir=tmp_path,
        output_filename="completed_real_1011.docx",
    )

    # 1. Verify success and file creation
    assert result.success is True
    assert result.completed_file.exists()
    assert result.completed_file.name == "completed_real_1011.docx"

    # 2. Verify original file untouched
    curr_hash = hashlib.sha256(real_path.read_bytes()).hexdigest()
    assert orig_hash == curr_hash

    # 3. Verify answers generated for both activities
    assert result.answers.total_count == 2
    for ans in result.answers.answers:
        assert ans.status == AnswerStatus.SUCCESS
        assert ans.question_type == QuestionType.LONG_ANSWER
        assert len(ans.answer_text) > 100

    # 4. Verify completed document structure and verification
    assert "verification" in result.summary
    assert result.summary["verification"]["is_valid"] is True
    comp_doc = docx.Document(str(result.completed_file))
    comp_text = "\n".join(p.text for p in comp_doc.paragraphs)
    assert len(comp_text) > 500
    assert "Architecture" in comp_text or "Technical" in comp_text or "Software" in comp_text


def test_default_worksheet_processor_adapter(tmp_path: Path):
    """Verify DefaultWorksheetProcessor processes solutions end-to-end."""
    orig_file = _create_synthetic_short_answer_docx(tmp_path / "proc_test.docx")
    processor = DefaultWorksheetProcessor()

    parsed_data = processor.parse_worksheet(orig_file)
    assert "questions" in parsed_data
    assert "document_path" in parsed_data

    completed_path = processor.process_solutions(parsed_data)
    assert completed_path.exists()
    assert completed_path != orig_file
    assert completed_path.name == "completed_proc_test.docx"
